import json
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any, ClassVar, cast
from unittest.mock import Mock, patch

from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.signing import TimestampSigner
from django.db import IntegrityError, connection, transaction
from django.test import Client, TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import Resolver404, resolve
from django.utils import timezone
from django.views.static import serve
from openpyxl import load_workbook
from rest_framework.test import APITestCase
from rest_framework_simplejwt.tokens import AccessToken, RefreshToken

from api.models import (
    BankSampah,
    BankSampahApprovalLog,
    JadwalKegiatan,
    JenisSampah,
    Nasabah,
    NasabahApprovalLog,
    Saldo,
    User,
)
from api.serializers import BankSampahApprovalListSerializer


class APISpecTests(APITestCase):
    _fake_tokens: ClassVar[Any]

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        cls._fake_tokens = override_settings(DEBUG=True, PILAH_ALLOW_FAKE_GOOGLE_TOKEN=True)
        cls._fake_tokens.enable()

    @classmethod
    def tearDownClass(cls) -> None:
        cls._fake_tokens.disable()
        super().tearDownClass()

    def setUp(self) -> None:
        self.bank = BankSampah.objects.create(
            nama="Bank Sampah BTH", alamat="Depok", kota="Depok", no_hp_pic="+628123456789"
        )
        self.user = User.objects.create_user(
            email="sari@example.com",
            nama="Ibu Sari",
            bank_sampah=self.bank,
            is_profile_complete=True,
            is_primary_pengelola=True,
        )
        refresh = RefreshToken.for_user(self.user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {refresh.access_token}")

    def test_bank_profile_update_normalizes_phone(self) -> None:
        response = self.client.put(
            "/api/v1/bank-sampah/me",
            {
                "nama": "Bank Sampah BTH",
                "alamat": "Kel. Kukusan",
                "kota": "Depok",
                "no_hp_pic": "081234567890",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["no_hp_pic"], "+6281234567890")
        self.assertEqual(response.data["pengelola"]["email"], "sari@example.com")

    def test_pengelola_can_create_and_update_organization_schedule(self) -> None:
        starts_at = timezone.now() + timedelta(days=2)
        create_response = self.client.post(
            "/api/v1/jadwal",
            {
                "jenis_kegiatan": "penimbangan",
                "mulai_pada": starts_at.isoformat(),
                "selesai_pada": (starts_at + timedelta(hours=2)).isoformat(),
                "lokasi": "Balai Warga RW 04",
                "keterangan": "Bawa sampah yang sudah dipilah",
                "cakupan_penerima": "semua_nasabah",
            },
            format="json",
        )

        self.assertEqual(create_response.status_code, 201)
        self.assertEqual(create_response.data["status"], "draft")
        self.assertEqual(create_response.data["bank_sampah_id"], str(self.bank.id))
        self.assertFalse(create_response.data["peringatan_jadwal_bertumpuk"])

        update_response = self.client.patch(
            f"/api/v1/jadwal/{create_response.data['id']}",
            {"lokasi": "Kantor Bank Sampah BTH"},
            format="json",
        )

        self.assertEqual(update_response.status_code, 200)
        self.assertEqual(update_response.data["lokasi"], "Kantor Bank Sampah BTH")

    def test_schedule_overlap_warns_without_blocking_creation(self) -> None:
        starts_at = timezone.now() + timedelta(days=3)
        payload = {
            "jenis_kegiatan": "penimbangan",
            "mulai_pada": starts_at.isoformat(),
            "selesai_pada": (starts_at + timedelta(hours=2)).isoformat(),
            "lokasi": "Balai Warga RW 04",
            "keterangan": "",
            "cakupan_penerima": "semua_nasabah",
        }
        first = self.client.post("/api/v1/jadwal", payload, format="json")
        second = self.client.post(
            "/api/v1/jadwal",
            {
                **payload,
                "mulai_pada": (starts_at + timedelta(minutes=30)).isoformat(),
                "selesai_pada": (starts_at + timedelta(hours=1)).isoformat(),
            },
            format="json",
        )

        self.assertEqual(first.status_code, 201)
        self.assertEqual(second.status_code, 201)
        self.assertTrue(second.data["peringatan_jadwal_bertumpuk"])

    def test_schedule_requires_end_after_start_on_create_and_update(self) -> None:
        starts_at = timezone.now() + timedelta(days=3)
        payload = {
            "jenis_kegiatan": "penimbangan",
            "mulai_pada": (starts_at + timedelta(hours=3)).isoformat(),
            "selesai_pada": (starts_at + timedelta(hours=2)).isoformat(),
            "lokasi": "Balai Warga",
            "cakupan_penerima": "semua_nasabah",
        }
        invalid_create = self.client.post("/api/v1/jadwal", payload, format="json")

        self.assertEqual(invalid_create.status_code, 422)
        self.assertEqual(
            invalid_create.data["errors"]["selesai_pada"],
            ["Waktu selesai harus setelah waktu mulai"],
        )

        payload["mulai_pada"] = starts_at.isoformat()
        created = self.client.post("/api/v1/jadwal", payload, format="json")
        invalid_update = self.client.patch(
            f"/api/v1/jadwal/{created.data['id']}",
            {"mulai_pada": (starts_at + timedelta(hours=3)).isoformat()},
            format="json",
        )

        self.assertEqual(created.status_code, 201)
        self.assertEqual(invalid_update.status_code, 422)
        self.assertEqual(
            invalid_update.data["errors"]["selesai_pada"],
            ["Waktu selesai harus setelah waktu mulai"],
        )
        self.assertEqual(
            JadwalKegiatan.objects.get(pk=created.data["id"]).mulai_pada,
            starts_at,
        )

    def test_schedule_update_recomputes_overlap_warning(self) -> None:
        starts_at = timezone.now() + timedelta(days=3)
        JadwalKegiatan.objects.create(
            bank_sampah=self.bank,
            dibuat_oleh=self.user,
            jenis_kegiatan=JadwalKegiatan.JenisKegiatan.PENIMBANGAN,
            mulai_pada=starts_at,
            selesai_pada=starts_at + timedelta(hours=2),
            lokasi="Balai Warga",
        )
        moving_schedule = JadwalKegiatan.objects.create(
            bank_sampah=self.bank,
            dibuat_oleh=self.user,
            jenis_kegiatan=JadwalKegiatan.JenisKegiatan.PENIMBANGAN,
            mulai_pada=starts_at,
            selesai_pada=starts_at + timedelta(hours=2),
            lokasi="Kantor",
        )

        overlapping = self.client.patch(
            f"/api/v1/jadwal/{moving_schedule.id}", {"lokasi": "Balai Warga"}, format="json"
        )
        self.assertEqual(overlapping.status_code, 200)
        self.assertTrue(overlapping.data["peringatan_jadwal_bertumpuk"])

        separated = self.client.patch(
            f"/api/v1/jadwal/{moving_schedule.id}", {"lokasi": "Kantor"}, format="json"
        )
        self.assertEqual(separated.status_code, 200)
        self.assertFalse(separated.data["peringatan_jadwal_bertumpuk"])

    def test_schedule_list_overlap_warnings_do_not_add_queries_per_row(self) -> None:
        starts_at = timezone.now() + timedelta(days=3)

        def create_schedule(index: int) -> None:
            begins = starts_at + timedelta(minutes=index)
            JadwalKegiatan.objects.create(
                bank_sampah=self.bank,
                dibuat_oleh=self.user,
                jenis_kegiatan=JadwalKegiatan.JenisKegiatan.PENIMBANGAN,
                mulai_pada=begins,
                selesai_pada=begins + timedelta(hours=2),
                lokasi="Balai Warga RW 04",
            )

        create_schedule(0)
        with CaptureQueriesContext(connection) as one_schedule_queries:
            one_schedule = self.client.get("/api/v1/jadwal?page_size=100")

        for index in range(1, 5):
            create_schedule(index)
        with CaptureQueriesContext(connection) as five_schedule_queries:
            five_schedules = self.client.get("/api/v1/jadwal?page_size=100")

        self.assertEqual(one_schedule.data["count"], 1)
        self.assertEqual(five_schedules.data["count"], 5)
        self.assertEqual(len(one_schedule_queries), len(five_schedule_queries))

    def test_schedule_list_filters_and_paginates_a_selected_day(self) -> None:
        starts_at = timezone.localtime() + timedelta(days=3)
        starts_at = starts_at.replace(hour=9, minute=0, second=0, microsecond=0)
        for index in range(3):
            begins = starts_at + timedelta(hours=index)
            JadwalKegiatan.objects.create(
                bank_sampah=self.bank,
                dibuat_oleh=self.user,
                jenis_kegiatan=JadwalKegiatan.JenisKegiatan.PENIMBANGAN,
                mulai_pada=begins,
                selesai_pada=begins + timedelta(minutes=30),
                lokasi=f"Lokasi {index}",
            )
        other_day = starts_at + timedelta(days=1)
        JadwalKegiatan.objects.create(
            bank_sampah=self.bank,
            dibuat_oleh=self.user,
            jenis_kegiatan=JadwalKegiatan.JenisKegiatan.PENIMBANGAN,
            mulai_pada=other_day,
            selesai_pada=other_day + timedelta(minutes=30),
            lokasi="Tanggal lain",
        )

        first_page = self.client.get(
            f"/api/v1/jadwal?date={timezone.localdate(starts_at)}&page_size=2"
        )
        second_page = self.client.get(
            f"/api/v1/jadwal?date={timezone.localdate(starts_at)}&page_size=2&page=2"
        )

        self.assertEqual(first_page.status_code, 200)
        self.assertEqual(first_page.data["count"], 3)
        self.assertEqual(len(first_page.data["results"]), 2)
        self.assertIsNotNone(first_page.data["next"])
        self.assertEqual(second_page.status_code, 200)
        self.assertEqual(len(second_page.data["results"]), 1)
        self.assertIsNone(second_page.data["next"])
        self.assertTrue(
            all(
                item["mulai_pada"].startswith(timezone.localdate(starts_at).isoformat())
                for item in first_page.data["results"] + second_page.data["results"]
            )
        )

        invalid_date = self.client.get("/api/v1/jadwal?date=not-a-date")
        self.assertEqual(invalid_date.status_code, 422)

    def test_schedule_calendar_dates_returns_unique_dates_for_a_range(self) -> None:
        starts_at = timezone.localtime() + timedelta(days=3)
        starts_at = starts_at.replace(hour=9, minute=0, second=0, microsecond=0)
        next_day = starts_at + timedelta(days=1)
        for index in range(2):
            begins = starts_at + timedelta(hours=index)
            JadwalKegiatan.objects.create(
                bank_sampah=self.bank,
                dibuat_oleh=self.user,
                jenis_kegiatan=JadwalKegiatan.JenisKegiatan.PENIMBANGAN,
                mulai_pada=begins,
                selesai_pada=begins + timedelta(minutes=30),
                lokasi=f"Lokasi {index}",
            )
        JadwalKegiatan.objects.create(
            bank_sampah=self.bank,
            dibuat_oleh=self.user,
            jenis_kegiatan=JadwalKegiatan.JenisKegiatan.PENIMBANGAN,
            mulai_pada=next_day,
            selesai_pada=next_day + timedelta(minutes=30),
            lokasi="Tanggal berikutnya",
        )
        start_date = timezone.localdate(starts_at)
        end_date = timezone.localdate(next_day)

        response = self.client.get(
            "/api/v1/jadwal/calendar-dates",
            {"start_date": start_date.isoformat(), "end_date": end_date.isoformat()},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.data["dates"], [start_date.isoformat(), end_date.isoformat()]
        )
        invalid_range = self.client.get(
            "/api/v1/jadwal/calendar-dates",
            {"start_date": end_date.isoformat(), "end_date": start_date.isoformat()},
        )
        self.assertEqual(invalid_range.status_code, 422)

    def test_schedule_rejects_unapproved_recipients(self) -> None:
        recipient = self._create_pending_nasabah()
        starts_at = timezone.now() + timedelta(days=3)

        response = self.client.post(
            "/api/v1/jadwal",
            {
                "jenis_kegiatan": "penimbangan",
                "mulai_pada": starts_at.isoformat(),
                "selesai_pada": (starts_at + timedelta(hours=2)).isoformat(),
                "lokasi": "Balai Warga RW 04",
                "cakupan_penerima": "nasabah_terpilih",
                "penerima_ids": [str(recipient.id)],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 422)
        self.assertIn("penerima_ids", response.data["errors"])

    def test_schedule_rejects_recipients_from_another_bank(self) -> None:
        other_bank = BankSampah.objects.create(
            nama="Bank Sampah Lain",
            alamat="Bogor",
            kota="Bogor",
            no_hp_pic="+6281234567891",
        )
        other_user = User.objects.create_user(
            email="nasabah-other-bank@example.com",
            nama="Nasabah Bank Lain",
            role=User.Role.NASABAH,
            is_profile_complete=False,
        )
        other_recipient = Nasabah.objects.create(
            user=other_user,
            bank_sampah=other_bank,
            nomor="NAS-1000",
            nama=other_user.nama,
            alamat="Jl. Kenanga",
            no_hp="+628123451000",
            status=Nasabah.Status.APPROVED,
        )
        starts_at = timezone.now() + timedelta(days=3)
        response = self.client.post(
            "/api/v1/jadwal",
            {
                "jenis_kegiatan": "penimbangan",
                "mulai_pada": starts_at.isoformat(),
                "selesai_pada": (starts_at + timedelta(hours=2)).isoformat(),
                "lokasi": "Balai Warga RW 04",
                "cakupan_penerima": "nasabah_terpilih",
                "penerima_ids": [str(other_recipient.id)],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 422)
        self.assertIn("penerima_ids", response.data["errors"])

    def test_schedule_rejects_explicitly_clearing_selected_recipients(self) -> None:
        recipient = self._create_pending_nasabah()
        recipient.status = Nasabah.Status.APPROVED
        recipient.save(update_fields=["status"])
        starts_at = timezone.now() + timedelta(days=3)
        created = self.client.post(
            "/api/v1/jadwal",
            {
                "jenis_kegiatan": "penimbangan",
                "mulai_pada": starts_at.isoformat(),
                "selesai_pada": (starts_at + timedelta(hours=2)).isoformat(),
                "lokasi": "Balai Warga RW 04",
                "cakupan_penerima": "nasabah_terpilih",
                "penerima_ids": [str(recipient.id)],
            },
            format="json",
        )
        response = self.client.patch(
            f"/api/v1/jadwal/{created.data['id']}", {"penerima_ids": []}, format="json"
        )

        self.assertEqual(created.status_code, 201)
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.data["errors"]["penerima_ids"], ["Pilih minimal satu nasabah"])
        self.assertTrue(
            self.client.get(f"/api/v1/jadwal/{created.data['id']}").data["penerima_ids"]
        )

    def test_pengelola_can_publish_cancel_and_complete_schedules(self) -> None:
        starts_at = timezone.now() + timedelta(days=4)

        def create_schedule() -> str:
            response = self.client.post(
                "/api/v1/jadwal",
                {
                    "jenis_kegiatan": "penimbangan",
                    "mulai_pada": starts_at.isoformat(),
                    "selesai_pada": (starts_at + timedelta(hours=2)).isoformat(),
                    "lokasi": "Balai Warga RW 04",
                    "keterangan": "",
                    "cakupan_penerima": "semua_nasabah",
                },
                format="json",
            )
            self.assertEqual(response.status_code, 201)
            return str(response.data["id"])

        cancelled_id = create_schedule()
        published = self.client.post(f"/api/v1/jadwal/{cancelled_id}/terbitkan")
        cancelled = self.client.post(f"/api/v1/jadwal/{cancelled_id}/batalkan")
        invalid_completion = self.client.post(f"/api/v1/jadwal/{cancelled_id}/selesaikan")

        self.assertEqual(published.status_code, 200)
        self.assertEqual(published.data["status"], "diterbitkan")
        self.assertEqual(cancelled.status_code, 200)
        self.assertEqual(cancelled.data["status"], "dibatalkan")
        self.assertEqual(invalid_completion.status_code, 400)

        completed_id = create_schedule()
        self.client.post(f"/api/v1/jadwal/{completed_id}/terbitkan")
        completed = self.client.post(f"/api/v1/jadwal/{completed_id}/selesaikan")

        self.assertEqual(completed.status_code, 200)
        self.assertEqual(completed.data["status"], "selesai")

    def test_nasabah_only_sees_active_published_schedules_for_their_memberships(self) -> None:
        nasabah_user = User.objects.create_user(
            email="nasabah@example.com",
            nama="Budi Nasabah",
            role=User.Role.NASABAH,
            is_profile_complete=False,
        )
        membership = Nasabah.objects.create(
            user=nasabah_user,
            bank_sampah=self.bank,
            nomor="NAS-0100",
            nama="Budi Nasabah",
            alamat="Jl. Kenanga No. 10",
            no_hp="+628123450100",
            status=Nasabah.Status.APPROVED,
        )
        another_bank = BankSampah.objects.create(
            nama="Bank Lain", alamat="Bogor", kota="Bogor", no_hp_pic="+628123456700"
        )
        now = timezone.now()

        visible = JadwalKegiatan.objects.create(
            bank_sampah=self.bank,
            dibuat_oleh=self.user,
            jenis_kegiatan=JadwalKegiatan.JenisKegiatan.PENIMBANGAN,
            mulai_pada=now + timedelta(days=1),
            selesai_pada=now + timedelta(days=1, hours=2),
            lokasi="Balai Warga",
            status=JadwalKegiatan.Status.DITERBITKAN,
        )
        selected = JadwalKegiatan.objects.create(
            bank_sampah=self.bank,
            dibuat_oleh=self.user,
            jenis_kegiatan=JadwalKegiatan.JenisKegiatan.PENCAIRAN,
            mulai_pada=now + timedelta(days=2),
            selesai_pada=now + timedelta(days=2, hours=1),
            lokasi="Kantor BTH",
            cakupan_penerima=JadwalKegiatan.CakupanPenerima.NASABAH_TERPILIH,
            status=JadwalKegiatan.Status.DITERBITKAN,
        )
        selected.penerima.add(membership)
        JadwalKegiatan.objects.create(
            bank_sampah=self.bank,
            dibuat_oleh=self.user,
            jenis_kegiatan=JadwalKegiatan.JenisKegiatan.PENIMBANGAN,
            mulai_pada=now + timedelta(days=3),
            selesai_pada=now + timedelta(days=3, hours=1),
            lokasi="Masih Draft",
        )
        JadwalKegiatan.objects.create(
            bank_sampah=another_bank,
            dibuat_oleh=self.user,
            jenis_kegiatan=JadwalKegiatan.JenisKegiatan.PENIMBANGAN,
            mulai_pada=now + timedelta(days=4),
            selesai_pada=now + timedelta(days=4, hours=1),
            lokasi="Organisasi Lain",
            status=JadwalKegiatan.Status.DITERBITKAN,
        )

        refresh = RefreshToken.for_user(nasabah_user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {refresh.access_token}")
        response = self.client.get("/api/v1/jadwal")

        self.assertEqual(response.status_code, 200)
        ids = {item["id"] for item in response.data["results"]}
        self.assertEqual(ids, {str(visible.id), str(selected.id)})

    def test_transition_does_not_overwrite_a_concurrent_status_change(self) -> None:
        from api.views import JadwalKegiatanViewSet

        starts_at = timezone.now() + timedelta(days=4)
        created = self.client.post(
            "/api/v1/jadwal",
            {
                "jenis_kegiatan": "penimbangan",
                "mulai_pada": starts_at.isoformat(),
                "selesai_pada": (starts_at + timedelta(hours=2)).isoformat(),
                "lokasi": "Balai Warga RW 04",
                "cakupan_penerima": "semua_nasabah",
            },
            format="json",
        )
        original_get_object = JadwalKegiatanViewSet.get_object

        def change_state_after_read(view: Any) -> Any:
            schedule = original_get_object(view)
            JadwalKegiatan.objects.filter(pk=schedule.pk).update(
                status=JadwalKegiatan.Status.DIBATALKAN
            )
            return schedule

        with patch.object(JadwalKegiatanViewSet, "get_object", change_state_after_read):
            response = self.client.post(f"/api/v1/jadwal/{created.data['id']}/terbitkan")

        self.assertEqual(response.status_code, 400)
        self.assertEqual(
            self.client.get(f"/api/v1/jadwal/{created.data['id']}").data["status"], "dibatalkan"
        )

    def test_transition_response_uses_fresh_schedule_after_concurrent_edit(self) -> None:
        from api.views import JadwalKegiatanViewSet

        recipient_before = self._create_pending_nasabah()
        recipient_before.status = Nasabah.Status.APPROVED
        recipient_before.save(update_fields=["status"])
        recipient_after = self._create_pending_nasabah()
        recipient_after.status = Nasabah.Status.APPROVED
        recipient_after.save(update_fields=["status"])
        starts_at = timezone.now() + timedelta(days=4)
        JadwalKegiatan.objects.create(
            bank_sampah=self.bank,
            dibuat_oleh=self.user,
            jenis_kegiatan=JadwalKegiatan.JenisKegiatan.PENIMBANGAN,
            mulai_pada=starts_at,
            selesai_pada=starts_at + timedelta(hours=2),
            lokasi="Balai Warga Baru",
        )
        schedule = JadwalKegiatan.objects.create(
            bank_sampah=self.bank,
            dibuat_oleh=self.user,
            jenis_kegiatan=JadwalKegiatan.JenisKegiatan.PENIMBANGAN,
            mulai_pada=starts_at,
            selesai_pada=starts_at + timedelta(hours=2),
            lokasi="Balai Warga Lama",
            cakupan_penerima=JadwalKegiatan.CakupanPenerima.NASABAH_TERPILIH,
        )
        schedule.penerima.add(recipient_before)
        original_get_object = JadwalKegiatanViewSet.get_object
        changed_after_read = False

        def edit_after_read(view: Any) -> Any:
            nonlocal changed_after_read
            instance = original_get_object(view)
            if not changed_after_read:
                changed_after_read = True
                JadwalKegiatan.objects.filter(pk=instance.pk).update(lokasi="Balai Warga Baru")
                through = JadwalKegiatan.penerima.through
                through.objects.filter(jadwalkegiatan_id=instance.pk).delete()
                through.objects.create(
                    jadwalkegiatan_id=instance.pk,
                    nasabah_id=recipient_after.pk,
                )
            return instance

        with patch.object(JadwalKegiatanViewSet, "get_object", edit_after_read):
            response = self.client.post(f"/api/v1/jadwal/{schedule.id}/batalkan")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["status"], "dibatalkan")
        self.assertEqual(response.data["lokasi"], "Balai Warga Baru")
        self.assertEqual(response.data["penerima_ids"], [recipient_after.id])
        self.assertTrue(response.data["peringatan_jadwal_bertumpuk"])

    def test_edit_does_not_overwrite_a_concurrent_terminal_transition(self) -> None:
        from api.views import JadwalKegiatanViewSet

        starts_at = timezone.now() + timedelta(days=4)
        created = self.client.post(
            "/api/v1/jadwal",
            {
                "jenis_kegiatan": "penimbangan",
                "mulai_pada": starts_at.isoformat(),
                "selesai_pada": (starts_at + timedelta(hours=2)).isoformat(),
                "lokasi": "Balai Warga RW 04",
                "cakupan_penerima": "semua_nasabah",
            },
            format="json",
        )
        original_get_object = JadwalKegiatanViewSet.get_object

        def cancel_after_read(view: Any) -> Any:
            schedule = original_get_object(view)
            JadwalKegiatan.objects.filter(pk=schedule.pk).update(
                status=JadwalKegiatan.Status.DIBATALKAN
            )
            return schedule

        with patch.object(JadwalKegiatanViewSet, "get_object", cancel_after_read):
            response = self.client.patch(
                f"/api/v1/jadwal/{created.data['id']}",
                {"lokasi": "Lokasi Baru"},
                format="json",
            )

        self.assertEqual(response.status_code, 400)
        schedule = JadwalKegiatan.objects.get(pk=created.data["id"])
        self.assertEqual(schedule.status, JadwalKegiatan.Status.DIBATALKAN)
        self.assertEqual(schedule.lokasi, "Balai Warga RW 04")

    def test_terminal_schedules_cannot_be_edited(self) -> None:
        starts_at = timezone.now() + timedelta(days=4)

        def create_schedule() -> str:
            response = self.client.post(
                "/api/v1/jadwal",
                {
                    "jenis_kegiatan": "penimbangan",
                    "mulai_pada": starts_at.isoformat(),
                    "selesai_pada": (starts_at + timedelta(hours=2)).isoformat(),
                    "lokasi": "Balai Warga RW 04",
                    "cakupan_penerima": "semua_nasabah",
                },
                format="json",
            )
            self.assertEqual(response.status_code, 201)
            return str(response.data["id"])

        cancelled_id = create_schedule()
        self.client.post(f"/api/v1/jadwal/{cancelled_id}/batalkan")
        completed_id = create_schedule()
        self.client.post(f"/api/v1/jadwal/{completed_id}/terbitkan")
        self.client.post(f"/api/v1/jadwal/{completed_id}/selesaikan")

        for schedule_id in (cancelled_id, completed_id):
            response = self.client.patch(
                f"/api/v1/jadwal/{schedule_id}", {"lokasi": "Lokasi Baru"}, format="json"
            )

            self.assertEqual(response.status_code, 400)
            schedule = self.client.get(f"/api/v1/jadwal/{schedule_id}")
            self.assertEqual(schedule.data["lokasi"], "Balai Warga RW 04")

    @override_settings(
        ALLOWED_HOSTS=["admin.example.com"],
        CSRF_TRUSTED_ORIGINS=["https://admin.example.com"],
        SECURE_PROXY_SSL_HEADER=("HTTP_X_FORWARDED_PROTO", "https"),
        STORAGES={
            "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
            "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
        },
    )
    def test_admin_login_accepts_cloud_run_https_origin(self) -> None:
        client = Client(enforce_csrf_checks=True)
        request_headers: dict[str, str] = {
            "HTTP_HOST": "admin.example.com",
            "HTTP_X_FORWARDED_PROTO": "https",
        }
        login_page = client.get("/admin/login/", **cast(dict[str, Any], request_headers))
        csrf_token = login_page.cookies["csrftoken"].value

        response = client.post(
            "/admin/login/",
            {
                "username": "missing@example.com",
                "password": "invalid",
                "csrfmiddlewaretoken": csrf_token,
            },
            HTTP_ORIGIN="https://admin.example.com",
            HTTP_REFERER="https://admin.example.com/admin/login/",
            **cast(dict[str, Any], request_headers),
        )

        self.assertEqual(response.status_code, 200)

    def test_bank_profile_update_accepts_logo_upload(self) -> None:
        response = self.client.put(
            "/api/v1/bank-sampah/me",
            {
                "nama": "Bank Sampah BTH",
                "alamat": "Kel. Kukusan",
                "kota": "Depok",
                "no_hp_pic": "081234567890",
                "foto_logo": SimpleUploadedFile("logo.png", b"logo", content_type="image/png"),
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("bank_sampah/logo/", response.data["foto_logo"])
        self.bank.refresh_from_db()
        assert self.bank.foto_logo.name is not None
        self.assertTrue(self.bank.foto_logo.name.startswith("bank_sampah/logo/"))

    def test_nasabah_create_list_detail_status_and_saldo(self) -> None:
        created = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NAS-0001",
                "nama": "Budi Santoso",
                "jenis_kelamin": "laki-laki",
                "tanggal_lahir": "1990-01-01",
                "no_hp": "081234567890",
                "alamat": "Jl. Anggrek No. 3",
            },
            format="json",
        )

        self.assertEqual(created.status_code, 201)
        self.assertEqual(created.data["kode"], "NAS-0001")
        self.assertEqual(created.data["total_saldo"], Decimal("0.00"))

        listed = self.client.get("/api/v1/nasabah?status=aktif")
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(listed.data["count"], 1)

        detail = self.client.get(f"/api/v1/nasabah/{created.data['id']}")
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.data["ringkasan_transaksi"]["jumlah_transaksi"], 0)

        patched = self.client.patch(
            f"/api/v1/nasabah/{created.data['id']}/status", {"is_active": False}, format="json"
        )
        self.assertEqual(patched.status_code, 200)
        self.assertFalse(patched.data["is_active"])

        inactive_update = self.client.put(
            f"/api/v1/nasabah/{created.data['id']}",
            {
                "kode": "NAS-0001",
                "nama": "Budi Santoso",
                "jenis_kelamin": "laki-laki",
                "tanggal_lahir": "1990-01-01",
                "no_hp": "081234567890",
                "alamat": "Jl. Anggrek No. 3",
            },
            format="json",
        )
        self.assertEqual(inactive_update.status_code, 403)

    def test_nasabah_created_by_pengurus_is_approved(self) -> None:
        created = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NAS-0001",
                "nama": "Budi Santoso",
                "jenis_kelamin": "laki-laki",
                "tanggal_lahir": "1990-01-01",
                "no_hp": "081234567890",
                "alamat": "Jl. Anggrek No. 3",
            },
            format="json",
        )

        self.assertEqual(created.status_code, 201)
        self.assertEqual(created.data["status"], "approved")
        self.nasabah = Nasabah.objects.get(id=created.data["id"])
        self.assertEqual(self.nasabah.status, Nasabah.Status.APPROVED)

    def _create_pending_nasabah(self, kode: str = "NAS-0001") -> Nasabah:
        no_hp = "081234567891"
        counter = getattr(self, "_pending_counter", 0)
        if counter:
            no_hp = f"08123456789{counter + 1}"
            kode = f"NAS-{counter + 1:04d}"
        self._pending_counter = counter + 1
        return Nasabah.objects.create(
            bank_sampah=self.bank,
            nomor=kode,
            nama="Citra Lestari",
            jenis_kelamin="perempuan",
            alamat="Jl. Melati No. 7",
            no_hp=no_hp,
            is_active=True,
            status=Nasabah.Status.PENDING,
        )

    def test_nasabah_list_status_filter_includes_menunggu_and_ditolak(self) -> None:
        pending = self._create_pending_nasabah("NAS-0001")
        rejected = Nasabah.objects.create(
            bank_sampah=self.bank,
            nomor="NAS-0002",
            nama="Dedi Pratama",
            jenis_kelamin="laki-laki",
            alamat="Jl. Kenanga No. 2",
            no_hp="081234567892",
            status=Nasabah.Status.REJECTED,
            is_active=False,
        )

        menunggu = self.client.get("/api/v1/nasabah?status=menunggu")
        self.assertEqual(menunggu.status_code, 200)
        self.assertEqual(menunggu.data["count"], 1)
        self.assertEqual(str(menunggu.data["results"][0]["id"]), str(pending.id))

        aktif = self.client.get("/api/v1/nasabah?status=aktif")
        self.assertEqual(aktif.data["count"], 0)

        ditolak = self.client.get("/api/v1/nasabah?status=ditolak")
        self.assertEqual(ditolak.status_code, 200)
        self.assertEqual(ditolak.data["count"], 1)
        self.assertEqual(str(ditolak.data["results"][0]["id"]), str(rejected.id))

        semua = self.client.get("/api/v1/nasabah?status=semua")
        self.assertEqual(semua.data["count"], 2)

    def test_nasabah_approve_pending_sets_approved_and_active(self) -> None:
        pending = self._create_pending_nasabah()

        response = self.client.post(
            f"/api/v1/nasabah/{pending.id}/approve", {"catatan": "Data lengkap"}, format="json"
        )

        self.assertEqual(response.status_code, 200)
        pending.refresh_from_db()
        self.assertEqual(pending.status, Nasabah.Status.APPROVED)
        self.assertTrue(pending.is_active)
        log = pending.approval_logs.get()
        self.assertEqual(log.status, NasabahApprovalLog.Status.APPROVED)
        self.assertEqual(log.catatan, "Data lengkap")
        self.assertEqual(log.pengurus, self.user)
        self.assertEqual(response.data["approval_log"]["status"], "approved")

    def test_nasabah_reject_pending_keeps_row_with_alasan(self) -> None:
        pending = self._create_pending_nasabah()

        response = self.client.post(
            f"/api/v1/nasabah/{pending.id}/reject", {"catatan": "Alamat tidak valid"}, format="json"
        )

        self.assertEqual(response.status_code, 200)
        pending.refresh_from_db()
        self.assertEqual(pending.status, Nasabah.Status.REJECTED)
        self.assertFalse(pending.is_active)
        log = pending.approval_logs.get()
        self.assertEqual(log.status, NasabahApprovalLog.Status.REJECTED)
        self.assertEqual(log.catatan, "Alamat tidak valid")
        self.assertEqual(log.pengurus, self.user)

    def test_nasabah_approve_non_pending_returns_400(self) -> None:
        approved = self._create_pending_nasabah()
        approved.status = Nasabah.Status.APPROVED
        approved.save(update_fields=["status"])

        response = self.client.post(f"/api/v1/nasabah/{approved.id}/approve", format="json")

        self.assertEqual(response.status_code, 400)

        rejected = self._create_pending_nasabah()
        rejected.status = Nasabah.Status.REJECTED
        rejected.save(update_fields=["status"])

        reject_response = self.client.post(f"/api/v1/nasabah/{rejected.id}/reject", format="json")
        self.assertEqual(reject_response.status_code, 400)

    def test_nasabah_other_bank_pending_not_accessible(self) -> None:
        other_bank = BankSampah.objects.create(
            nama="Bank Sampah Lain", alamat="Bogor", kota="Bogor", no_hp_pic="+628129876543"
        )
        pending = Nasabah.objects.create(
            bank_sampah=other_bank,
            nomor="NAS-0001",
            nama="Eko Widodo",
            jenis_kelamin="laki-laki",
            alamat="Jl. Cendana No. 9",
            no_hp="081234567893",
            status=Nasabah.Status.PENDING,
        )

        response = self.client.post(f"/api/v1/nasabah/{pending.id}/approve", format="json")

        self.assertEqual(response.status_code, 404)

    def test_nasabah_duplicate_phone_returns_validation_error(self) -> None:
        first = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NAS-0001",
                "nama": "Budi Santoso",
                "jenis_kelamin": "laki-laki",
                "tanggal_lahir": "1990-01-01",
                "no_hp": "081234567890",
                "alamat": "Jl. Anggrek No. 3",
            },
            format="json",
        )
        self.assertEqual(first.status_code, 201)

        duplicate_create = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NAS-0002",
                "nama": "Dewi Lestari",
                "jenis_kelamin": "perempuan",
                "tanggal_lahir": "1992-02-02",
                "no_hp": "+6281234567890",
                "alamat": "Jl. Melati No. 7",
            },
            format="json",
        )
        self.assertEqual(duplicate_create.status_code, 422)
        self.assertEqual(
            duplicate_create.data["errors"]["no_hp"], ["Nomor HP nasabah sudah digunakan"]
        )

        second = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NAS-0002",
                "nama": "Dewi Lestari",
                "jenis_kelamin": "perempuan",
                "tanggal_lahir": "1992-02-02",
                "no_hp": "081999999999",
                "alamat": "Jl. Melati No. 7",
            },
            format="json",
        )
        self.assertEqual(second.status_code, 201)

        duplicate_update = self.client.put(
            f"/api/v1/nasabah/{second.data['id']}",
            {
                "kode": "NAS-0002",
                "nama": "Dewi Lestari",
                "jenis_kelamin": "perempuan",
                "tanggal_lahir": "1992-02-02",
                "no_hp": "081234567890",
                "alamat": "Jl. Melati No. 7",
            },
            format="json",
        )
        self.assertEqual(duplicate_update.status_code, 422)
        self.assertEqual(
            duplicate_update.data["errors"]["no_hp"], ["Nomor HP nasabah sudah digunakan"]
        )

    def test_nasabah_create_with_email(self) -> None:
        created = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NAS-0001",
                "nama": "Budi Santoso",
                "jenis_kelamin": "laki-laki",
                "tanggal_lahir": "1990-01-01",
                "no_hp": "081234567890",
                "alamat": "Jl. Anggrek No. 3",
                "email": "Budi@Example.com",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201)
        self.assertEqual(created.data["email"], "budi@example.com")
        self.assertTrue(Saldo.objects.filter(nasabah_id=created.data["id"]).exists())

    def test_nasabah_create_without_email_still_works(self) -> None:
        created = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NAS-0001",
                "nama": "Budi Santoso",
                "jenis_kelamin": "laki-laki",
                "tanggal_lahir": "1990-01-01",
                "no_hp": "081234567890",
                "alamat": "Jl. Anggrek No. 3",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201)
        self.assertIsNone(created.data["email"])

    def test_nasabah_duplicate_email_same_bank_returns_validation_error(self) -> None:
        first = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NAS-0001",
                "nama": "Budi Santoso",
                "jenis_kelamin": "laki-laki",
                "tanggal_lahir": "1990-01-01",
                "no_hp": "081234567890",
                "alamat": "Jl. Anggrek No. 3",
                "email": "budi@example.com",
            },
            format="json",
        )
        self.assertEqual(first.status_code, 201)

        duplicate = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NAS-0002",
                "nama": "Dewi Lestari",
                "jenis_kelamin": "perempuan",
                "tanggal_lahir": "1992-02-02",
                "no_hp": "081999999999",
                "alamat": "Jl. Melati No. 7",
                "email": "Budi@example.com",
            },
            format="json",
        )
        self.assertEqual(duplicate.status_code, 422)
        self.assertEqual(
            duplicate.data["errors"]["email"],
            ["Email sudah terdaftar sebagai nasabah di bank sampah ini"],
        )

    def test_nasabah_duplicate_email_other_bank_returns_validation_error(self) -> None:
        created = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NAS-0001",
                "nama": "Budi Santoso",
                "jenis_kelamin": "laki-laki",
                "tanggal_lahir": "1990-01-01",
                "no_hp": "081234567890",
                "alamat": "Jl. Anggrek No. 3",
                "email": "budi@example.com",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201)

        other_bank = BankSampah.objects.create(
            nama="Bank Sampah Lain", alamat="Jakarta", kota="Jakarta", no_hp_pic="+628123456788"
        )
        other_user = User.objects.create_user(
            email="pengelola-lain@example.com",
            nama="Pengelola Lain",
            bank_sampah=other_bank,
            is_profile_complete=True,
            is_primary_pengelola=True,
        )
        refresh = RefreshToken.for_user(other_user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {refresh.access_token}")

        duplicate = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NSL-0001",
                "nama": "Budi Santoso",
                "jenis_kelamin": "laki-laki",
                "tanggal_lahir": "1990-01-01",
                "no_hp": "081234567890",
                "alamat": "Jl. Anggrek No. 3",
                "email": "budi@example.com",
            },
            format="json",
        )
        self.assertEqual(duplicate.status_code, 422)
        self.assertEqual(
            duplicate.data["errors"]["email"],
            ["Email sudah terdaftar sebagai nasabah di bank sampah lain"],
        )

    def test_nasabah_email_on_update(self) -> None:
        first = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NAS-0001",
                "nama": "Budi Santoso",
                "jenis_kelamin": "laki-laki",
                "tanggal_lahir": "1990-01-01",
                "no_hp": "081234567890",
                "alamat": "Jl. Anggrek No. 3",
                "email": "budi@example.com",
            },
            format="json",
        )
        self.assertEqual(first.status_code, 201)

        second = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NAS-0002",
                "nama": "Dewi Lestari",
                "jenis_kelamin": "perempuan",
                "tanggal_lahir": "1992-02-02",
                "no_hp": "081999999999",
                "alamat": "Jl. Melati No. 7",
                "email": "dewi@example.com",
            },
            format="json",
        )
        self.assertEqual(second.status_code, 201)

        stolen = self.client.put(
            f"/api/v1/nasabah/{second.data['id']}",
            {
                "kode": "NAS-0002",
                "nama": "Dewi Lestari",
                "jenis_kelamin": "perempuan",
                "tanggal_lahir": "1992-02-02",
                "no_hp": "081999999999",
                "alamat": "Jl. Melati No. 7",
                "email": "budi@example.com",
            },
            format="json",
        )
        self.assertEqual(stolen.status_code, 422)
        self.assertEqual(
            stolen.data["errors"]["email"],
            ["Email sudah terdaftar sebagai nasabah di bank sampah ini"],
        )

        kept = self.client.put(
            f"/api/v1/nasabah/{second.data['id']}",
            {
                "kode": "NAS-0002",
                "nama": "Dewi Lestari",
                "jenis_kelamin": "perempuan",
                "tanggal_lahir": "1992-02-02",
                "no_hp": "081999999999",
                "alamat": "Jl. Melati No. 7",
                "email": "dewi@example.com",
            },
            format="json",
        )
        self.assertEqual(kept.status_code, 200)

    def test_nasabah_invalid_email_format_returns_validation_error(self) -> None:
        response = self.client.post(
            "/api/v1/nasabah",
            {
                "kode": "NAS-0001",
                "nama": "Budi Santoso",
                "jenis_kelamin": "laki-laki",
                "tanggal_lahir": "1990-01-01",
                "no_hp": "081234567890",
                "alamat": "Jl. Anggrek No. 3",
                "email": "bukan-email",
            },
            format="json",
        )
        self.assertEqual(response.status_code, 422)
        self.assertTrue(response.data["errors"]["email"])

    def test_google_login_with_nasabah_email_prefills_profile(self) -> None:
        self.client.credentials()
        Nasabah.objects.create(
            bank_sampah=self.bank,
            nomor="NAS-0001",
            nama="Budi Santoso",
            jenis_kelamin="laki-laki",
            tanggal_lahir="1990-01-01",
            alamat="Jl. Anggrek No. 3",
            no_hp="081234567890",
            email="budi@example.com",
        )

        response = self.client.post(
            "/api/v1/auth/google", {"id_token": "dev:budi@example.com:B"}, format="json"
        )
        self.assertEqual(response.status_code, 200)

        user = User.objects.get(email="budi@example.com")
        self.assertEqual(user.nama, "Budi Santoso")
        self.assertEqual(user.no_hp, "081234567890")
        self.assertTrue(user.is_profile_complete)
        self.assertEqual(response.data["next_step"], "register_bank_sampah")

    def test_google_login_does_not_overwrite_existing_profile(self) -> None:
        self.client.credentials()
        User.objects.create_user(
            email="lengkap@example.com",
            nama="Nama Lama",
            no_hp="081111111111",
            is_profile_complete=True,
        )
        Nasabah.objects.create(
            bank_sampah=self.bank,
            nomor="NAS-0001",
            nama="Nama Nasabah",
            alamat="Jl. Anggrek No. 3",
            no_hp="081234567890",
            email="lengkap@example.com",
        )

        response = self.client.post(
            "/api/v1/auth/google",
            {"id_token": "dev:lengkap@example.com:Baru"},
            format="json",
        )
        self.assertEqual(response.status_code, 200)

        user = User.objects.get(email="lengkap@example.com")
        self.assertEqual(user.nama, "Nama Lama")
        self.assertEqual(user.no_hp, "081111111111")

    def test_google_login_prefills_nama_from_nasabah_when_existing_name_empty(
        self,
    ) -> None:
        """P2 regression: the token fallback must not block the nasabah name."""
        self.client.credentials()
        User.objects.create_user(
            email="kosong@example.com",
            nama="",
            is_profile_complete=False,
        )
        Nasabah.objects.create(
            bank_sampah=self.bank,
            nomor="NAS-0001",
            nama="Nama Nasabah",
            alamat="Jl. Anggrek No. 3",
            no_hp="081234567890",
            email="kosong@example.com",
        )

        response = self.client.post(
            "/api/v1/auth/google", {"id_token": "dev:kosong@example.com:K"}, format="json"
        )
        self.assertEqual(response.status_code, 200)

        user = User.objects.get(email="kosong@example.com")
        self.assertEqual(user.nama, "Nama Nasabah")

    def test_google_login_partial_nasabah_keeps_profile_incomplete(self) -> None:
        """P1 regression: a nasabah record lacking jenis_kelamin/tanggal_lahir
        must not mark the profile complete — /onboarding/profile would reject
        the follow-up with 'Profil sudah lengkap'."""
        self.client.credentials()
        Nasabah.objects.create(
            bank_sampah=self.bank,
            nomor="NAS-0001",
            nama="Budi Sebagian",
            alamat="Jl. Anggrek No. 3",
            no_hp="081234567890",
            email="sebagian@example.com",
        )

        response = self.client.post(
            "/api/v1/auth/google",
            {"id_token": "dev:sebagian@example.com:S"},
            format="json",
        )
        self.assertEqual(response.status_code, 200)

        user = User.objects.get(email="sebagian@example.com")
        self.assertEqual(user.nama, "Budi Sebagian")
        self.assertEqual(user.no_hp, "081234567890")
        self.assertFalse(user.jenis_kelamin)
        self.assertIsNone(user.tanggal_lahir)
        self.assertFalse(user.is_profile_complete)
        self.assertEqual(response.data["next_step"], "complete_profile")

    def test_jenis_sampah_and_transaction_update_saldo(self) -> None:
        nasabah = Nasabah.objects.create(
            bank_sampah=self.bank,
            nomor="NAS-0001",
            nama="Ahmad Ridwan",
            no_hp="+628123456789",
            alamat="Jl. Mawar No. 12",
        )
        Saldo.objects.create(nasabah=nasabah)
        jenis = self.client.post(
            "/api/v1/jenis-sampah",
            {
                "kode": "PLS-001",
                "nama_sampah": "Plastik PET",
                "kategori": "plastik",
                "deskripsi": "Botol bening",
                "harga_per_kg": 3500,
            },
            format="json",
        )
        self.assertEqual(jenis.status_code, 201)

        transaksi = self.client.post(
            "/api/v1/transaksi",
            {
                "nasabah_id": str(nasabah.id),
                "items": [{"jenis_sampah_id": jenis.data["id"], "berat": "2.500"}],
                "catatan": "Setoran rutin",
            },
            format="json",
        )
        self.assertEqual(transaksi.status_code, 201)
        self.assertEqual(transaksi.data["total_nilai"], "8750.00")
        self.assertEqual(transaksi.data["saldo_setelah_transaksi"], Decimal("8750.00"))
        self.assertEqual(transaksi.data["items"][0]["harga_snapshot"], "3500.00")

        saldo = self.client.get(f"/api/v1/nasabah/{nasabah.id}/saldo")
        self.assertEqual(saldo.status_code, 200)
        self.assertEqual(saldo.data["total_saldo"], "8750.00")

        with self.settings(
            TWILIO_ACCOUNT_SID="",
            TWILIO_AUTH_TOKEN="",
            TWILIO_API_KEY_SID="",
            TWILIO_API_KEY_SECRET="",
            TWILIO_WHATSAPP_FROM="",
            TWILIO_MESSAGING_SERVICE_SID="",
            WHATSAPP_GATEWAY_URL="",
        ):
            notify = self.client.post(f"/api/v1/transaksi/{transaksi.data['id']}/notify-wa")
        self.assertEqual(notify.status_code, 400)
        self.assertEqual(notify.data["status_wa"], "gagal")
        self.assertEqual(notify.data["error"], "Konfigurasi WhatsApp/Twilio belum diisi")

        transaksi_kedua = self.client.post(
            "/api/v1/transaksi",
            {
                "nasabah_id": str(nasabah.id),
                "items": [{"jenis_sampah_id": jenis.data["id"], "berat": "1.000"}],
                "catatan": "Setoran kedua",
            },
            format="json",
        )
        self.assertEqual(transaksi_kedua.status_code, 201)
        self.assertEqual(transaksi_kedua.data["saldo_setelah_transaksi"], Decimal("12250.00"))
        detail_pertama = self.client.get(f"/api/v1/transaksi/{transaksi.data['id']}")
        self.assertEqual(detail_pertama.status_code, 200)
        self.assertEqual(detail_pertama.data["saldo_setelah_transaksi"], Decimal("8750.00"))
        detail_kedua = self.client.get(f"/api/v1/transaksi/{transaksi_kedua.data['id']}")
        self.assertEqual(detail_kedua.status_code, 200)
        self.assertEqual(detail_kedua.data["saldo_setelah_transaksi"], Decimal("12250.00"))

        export = self.client.get("/api/v1/transaksi/export?periode=bulan_ini")
        self.assertEqual(export.status_code, 200)
        self.assertEqual(
            export["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("PILAH_Laporan_", export["Content-Disposition"])
        workbook = load_workbook(BytesIO(export.content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Laporan", "Riwayat Transaksi"])
        summary_sheet = workbook["Laporan"]
        self.assertEqual(
            [summary_sheet.cell(1, col).value for col in range(1, 6)],
            ["Jenis", "Sampah", "Harga per kg", "Jumlah kg", "Total"],
        )
        self.assertEqual(summary_sheet["A2"].value, "plastik")
        self.assertEqual(summary_sheet["B2"].value, "Plastik PET")
        self.assertEqual(summary_sheet["C2"].value, 3500)
        self.assertEqual(summary_sheet["D2"].value, 3.5)
        self.assertEqual(summary_sheet["E2"].value, "=C2*D2")
        sheet = workbook["Riwayat Transaksi"]
        self.assertEqual(sheet.freeze_panes, "A5")
        self.assertEqual(sheet["A1"].value, "PILAH - Riwayat Transaksi")
        self.assertIn("Filter periode: Bulan Ini", sheet["A2"].value)
        self.assertEqual(
            [sheet.cell(4, col).value for col in range(1, 11)],
            [
                "No",
                "Tanggal",
                "Waktu",
                "Nama Nasabah",
                "ID Nasabah",
                "Jenis Sampah",
                "Berat (kg)",
                "Harga/kg (Rp)",
                "Subtotal (Rp)",
                "Saldo Setelah Transaksi (Rp)",
            ],
        )
        self.assertEqual(sheet["A5"].value, 1)
        self.assertEqual(sheet["D5"].value, "Ahmad Ridwan")
        self.assertEqual(sheet["E5"].value, "NAS-0001")
        self.assertEqual(sheet["F5"].value, "Plastik PET")
        self.assertEqual(sheet["G5"].value, 1)
        self.assertEqual(sheet["H5"].value, 3500)
        self.assertEqual(sheet["I5"].value, 3500)
        self.assertEqual(sheet["J5"].value, 12250)
        self.assertEqual(sheet["G6"].value, 2.5)
        self.assertEqual(sheet["I6"].value, 8750)
        self.assertEqual(sheet["J6"].value, 8750)

    @patch("api.services.requests.post")
    def test_transaction_notify_wa_uses_twilio(self, post: Mock) -> None:
        post.return_value = Mock(status_code=201, json=lambda: {"sid": "SM123"}, text="")
        nasabah = Nasabah.objects.create(
            bank_sampah=self.bank,
            nomor="NAS-0002",
            nama="Dewi Lestari",
            no_hp="+628111111111",
            alamat="Jl. Melati No. 9",
        )
        Saldo.objects.create(nasabah=nasabah)
        jenis = JenisSampah.objects.create(
            bank_sampah=self.bank,
            nomor="JS-0002",
            nama_sampah="Kardus",
            kategori="kertas",
            harga_per_kg=Decimal("2000"),
        )
        transaksi = self.client.post(
            "/api/v1/transaksi",
            {
                "nasabah_id": str(nasabah.id),
                "items": [{"jenis_sampah_id": str(jenis.id), "berat": "1.000"}],
            },
            format="json",
        )

        with self.settings(
            TWILIO_ACCOUNT_SID="ACxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxx",
            TWILIO_AUTH_TOKEN="",
            TWILIO_API_KEY_SID="SKxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxx",
            TWILIO_API_KEY_SECRET="secret",
            TWILIO_WHATSAPP_FROM="whatsapp:+14155238886",
            TWILIO_MESSAGING_SERVICE_SID="",
            TWILIO_CONTENT_SID="HXb1844641bcade1dafdfecdf5f6a4aefd",
            WHATSAPP_GATEWAY_URL="",
        ):
            notify = self.client.post(f"/api/v1/transaksi/{transaksi.data['id']}/notify-wa")

        self.assertEqual(notify.status_code, 200)
        self.assertEqual(notify.data["status_wa"], "terkirim")
        self.assertEqual(notify.data["provider"], "twilio")
        content_variables = post.call_args.kwargs["data"]["ContentVariables"]
        self.assertEqual(
            json.loads(content_variables),
            {"1": "Dewi Lestari", "2": "- Kardus 1 kg"},
        )
        post.assert_called_once_with(
            "https://api.twilio.com/2010-04-01/Accounts/ACxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxx/Messages.json",
            data={
                "To": "whatsapp:+628111111111",
                "ContentSid": "HXb1844641bcade1dafdfecdf5f6a4aefd",
                "ContentVariables": content_variables,
                "From": "whatsapp:+14155238886",
            },
            auth=("SKxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxx", "secret"),
            timeout=10,
        )

    def test_dashboard_and_wa_template(self) -> None:
        stats = self.client.get("/api/v1/dashboard/stats")
        self.assertEqual(stats.status_code, 200)
        self.assertEqual(stats.data["bank_sampah_nama"], "Bank Sampah BTH")

        template = self.client.get("/api/v1/pengaturan/wa-template")
        self.assertEqual(template.status_code, 200)
        self.assertIn("{Nama}", template.data["variabel_tersedia"])

        updated = self.client.put(
            "/api/v1/pengaturan/wa-template",
            {"template": "Halo {Nama}, saldo {Saldo}"},
            format="json",
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.data["message"], "Template berhasil disimpan")

    @override_settings(
        TWILIO_ACCOUNT_SID="",
        TWILIO_AUTH_TOKEN="",
        TWILIO_API_KEY_SID="",
        TWILIO_API_KEY_SECRET="",
        TWILIO_WHATSAPP_FROM="",
        TWILIO_MESSAGING_SERVICE_SID="",
        WHATSAPP_GATEWAY_URL="https://wa.example.test/send",
        WHATSAPP_GATEWAY_TOKEN="test-token",
    )
    @patch("api.services.requests.post")
    def test_wa_template_item_variables_match_sent_payload(self, post: Mock) -> None:
        post.return_value = Mock(status_code=200, text="")
        self.client.put(
            "/api/v1/pengaturan/wa-template",
            {
                "template": (
                    "Halo {Nama}\n{daftar_item}\nLengkap:\n{daftar_item_harga}\nSaldo {Saldo}"
                )
            },
            format="json",
        )
        nasabah = Nasabah.objects.create(
            bank_sampah=self.bank,
            nomor="NAS-0002",
            nama="Budi Santoso",
            no_hp="+628123456789",
            alamat="Jl. Melati No. 5",
        )
        Saldo.objects.create(nasabah=nasabah)
        kertas = JenisSampah.objects.create(
            bank_sampah=self.bank,
            nomor="JS-0003",
            nama_sampah="kertas hvs",
            kategori="kertas",
            harga_per_kg=Decimal("1500"),
        )
        botol = JenisSampah.objects.create(
            bank_sampah=self.bank,
            nomor="JS-0004",
            nama_sampah="botol kaca",
            kategori="kaca",
            harga_per_kg=Decimal("1000"),
        )
        transaksi = self.client.post(
            "/api/v1/transaksi",
            {
                "nasabah_id": str(nasabah.id),
                "items": [
                    {"jenis_sampah_id": str(kertas.id), "berat": "2.300"},
                    {"jenis_sampah_id": str(botol.id), "berat": "17.123"},
                ],
            },
            format="json",
        )
        self.assertEqual(transaksi.status_code, 201)

        notify = self.client.post(f"/api/v1/transaksi/{transaksi.data['id']}/notify-wa")

        self.assertEqual(notify.status_code, 200, notify.data)
        sent_payload = post.call_args.kwargs["json"]
        self.assertEqual(
            sent_payload["message"],
            (
                "Halo Budi Santoso\n"
                "- kertas hvs 2,3 kg\n"
                "- botol kaca 17,12 kg\n"
                "Lengkap:\n"
                "- kertas hvs 2,3 kg x Rp 1.500 = Rp 3.450\n"
                "- botol kaca 17,12 kg x Rp 1.000 = Rp 17.123\n"
                "Saldo Rp 20.573"
            ),
        )

        template = self.client.get("/api/v1/pengaturan/wa-template")
        self.assertIn("{daftar_item_harga}", template.data["variabel_tersedia"])
        self.assertNotIn("....", template.data["preview_contoh"])

    def test_google_dev_auth(self) -> None:
        response = self.client.post(
            "/api/v1/auth/google", {"id_token": "dev:new@example.com:New User"}, format="json"
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data["registration_required"])
        self.assertIn("registration_token", response.data)
        self.assertFalse(User.objects.filter(email="new@example.com").exists())

    @override_settings(PILAH_SUPERADMIN_EMAILS=("super@example.com",))
    def test_onboarding_superadmin_approval_and_invite_flow(self) -> None:
        self.client.credentials()
        login = self.client.post(
            "/api/v1/auth/google",
            {"id_token": "dev-pengelola:onboard@example.com:Onboard User"},
            format="json",
        )
        self.assertEqual(login.status_code, 200)
        token = login.data["access_token"]
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")

        blocked = self.client.get("/api/v1/dashboard/stats")
        self.assertEqual(blocked.status_code, 403)

        profile = self.client.put(
            "/api/v1/onboarding/profile",
            {
                "nama": "Onboard User",
                "jenis_kelamin": "laki-laki",
                "tanggal_lahir": "1990-01-01",
                "no_hp": "081234567892",
            },
            format="json",
        )
        self.assertEqual(profile.status_code, 200)
        self.assertEqual(profile.data["next_step"], "register_bank_sampah")

        registration = self.client.post(
            "/api/v1/onboarding/bank-sampah",
            {
                "nama": "Bank Sampah Pending",
                "alamat": "Jl. Pending No. 1, Kukusan",
                "kota": "Depok",
                "no_hp_pic": "081234567893",
                "foto_kegiatan": SimpleUploadedFile(
                    "pending.png", b"png", content_type="image/png"
                ),
            },
            format="multipart",
        )
        self.assertEqual(registration.status_code, 201)
        self.assertEqual(registration.data["status"], "pending")
        self.assertEqual(registration.data["next_step"], "approval_pending")

        pending_block = self.client.get("/api/v1/nasabah")
        self.assertEqual(pending_block.status_code, 403)

        superadmin = User.objects.create_user(
            email="super@example.com",
            nama="Super Admin",
            role=User.Role.SUPERADMIN,
            is_profile_complete=True,
        )
        super_refresh = RefreshToken.for_user(superadmin)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {super_refresh.access_token}")

        queue = self.client.get("/api/v1/superadmin/bank-sampah?status=pending")
        self.assertEqual(queue.status_code, 200)
        self.assertEqual(queue.data["count"], 1)

        approval = self.client.post(
            f"/api/v1/superadmin/bank-sampah/{registration.data['id']}/approve",
            {"catatan": "Valid"},
            format="json",
        )
        self.assertEqual(approval.status_code, 200)
        self.assertEqual(approval.data["bank_sampah"]["status"], "active")
        self.assertEqual(BankSampahApprovalLog.objects.count(), 1)

        refreshed = self.client.post(
            "/api/v1/auth/refresh", {"refresh_token": login.data["refresh_token"]}, format="json"
        )
        self.assertEqual(refreshed.status_code, 200)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {refreshed.data['access_token']}")

        active_dashboard = self.client.get("/api/v1/dashboard/stats")
        self.assertEqual(active_dashboard.status_code, 200)

        profile_overwrite = self.client.put(
            "/api/v1/onboarding/profile",
            {
                "nama": "Overwritten User",
                "jenis_kelamin": "perempuan",
                "tanggal_lahir": "1999-09-09",
                "no_hp": "081299999999",
            },
            format="json",
        )
        self.assertEqual(profile_overwrite.status_code, 400)
        self.assertEqual(profile_overwrite.data["error"], "Profil sudah lengkap")

        invite = self.client.post("/api/v1/team/invite")
        self.assertEqual(invite.status_code, 201)
        self.assertIn("token", invite.data)
        expires_at = invite.data["expires_at"]
        self.assertGreaterEqual(
            expires_at, timezone.now() + timedelta(days=3) - timedelta(seconds=5)
        )
        self.assertLessEqual(expires_at, timezone.now() + timedelta(days=3) + timedelta(seconds=5))

        self_accept = self.client.post(
            "/api/v1/invites/accept", {"token": invite.data["token"]}, format="json"
        )
        self.assertEqual(self_accept.status_code, 200)
        self.assertEqual(self_accept.data["outcome"], "already_member")
        self.assertEqual(self_accept.data["message"], "Anda sudah terdaftar pada bank sampah ini")
        onboard_user = User.objects.get(email="onboard@example.com")
        self.assertTrue(onboard_user.is_primary_pengelola)
        self.assertEqual(str(onboard_user.bank_sampah_id), registration.data["id"])

        self.client.credentials()
        invited_login = self.client.post(
            "/api/v1/auth/google",
            {"id_token": "dev-pengelola:invited@example.com:Invited User"},
            format="json",
        )
        self.assertEqual(invited_login.status_code, 200)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {invited_login.data['access_token']}")
        invited_profile = self.client.put(
            "/api/v1/onboarding/profile",
            {
                "nama": "Invited User",
                "jenis_kelamin": "perempuan",
                "tanggal_lahir": "1992-02-02",
                "no_hp": "081234567894",
            },
            format="json",
        )
        self.assertEqual(invited_profile.status_code, 200)
        accepted = self.client.post(
            "/api/v1/invites/accept", {"token": invite.data["token"]}, format="json"
        )
        self.assertEqual(accepted.status_code, 200)
        self.assertEqual(accepted.data["next_step"], "dashboard")
        self.assertEqual(accepted.data["outcome"], "join_success")
        self.assertEqual(accepted.data["bank_sampah_id"], registration.data["id"])
        self.assertEqual(accepted.data["bank_sampah_nama"], "Bank Sampah Pending")

        repeated_accept = self.client.post(
            "/api/v1/invites/accept", {"token": invite.data["token"]}, format="json"
        )
        self.assertEqual(repeated_accept.status_code, 200)
        self.assertEqual(repeated_accept.data["outcome"], "already_member")

        team = self.client.get("/api/v1/team")
        self.assertEqual(team.status_code, 200)
        self.assertEqual(len(team.data["members"]), 2)

        denied_invite = self.client.post("/api/v1/team/invite")
        self.assertEqual(denied_invite.status_code, 403)

    def test_invite_join_rejects_superadmin_at_join_endpoint(self) -> None:
        active_bank = BankSampah.objects.create(
            nama="Invite Target", alamat="Depok", no_hp_pic="+628111111111"
        )
        invite_token = active_bank.invite_token = "target-token"
        active_bank.invite_token_expires = timezone.now() + timedelta(days=3)
        active_bank.save(update_fields=["invite_token", "invite_token_expires", "updated_at"])
        superadmin = User.objects.create_user(
            email="super-invite@example.com",
            nama="Super Invite",
            role=User.Role.SUPERADMIN,
            is_profile_complete=True,
        )
        refresh = RefreshToken.for_user(superadmin)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {refresh.access_token}")

        response = self.client.post(
            "/api/v1/bank-sampah/invite/join", {"invite_token": invite_token}, format="json"
        )

        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.data["error"], "Hanya pengelola yang dapat menerima undangan")

    def test_invite_join_rejects_active_and_pending_other_bank(self) -> None:
        target_bank = BankSampah.objects.create(
            nama="Target Active", alamat="Depok", no_hp_pic="+628111111111"
        )
        invite_token = target_bank.invite_token = "target-active-token"
        target_bank.invite_token_expires = timezone.now() + timedelta(days=3)
        target_bank.save(update_fields=["invite_token", "invite_token_expires", "updated_at"])
        active_bank = BankSampah.objects.create(
            nama="Current Active", alamat="Depok", no_hp_pic="+628222222222"
        )
        pending_bank = BankSampah.objects.create(
            nama="Current Pending",
            alamat="Depok",
            no_hp_pic="+628333333333",
            status=BankSampah.Status.PENDING,
            is_active=False,
        )

        active_user = User.objects.create_user(
            email="active-current@example.com",
            nama="Active Current",
            bank_sampah=active_bank,
            is_profile_complete=True,
            is_primary_pengelola=True,
        )
        pending_user = User.objects.create_user(
            email="pending-current@example.com",
            nama="Pending Current",
            bank_sampah=pending_bank,
            is_profile_complete=True,
            is_primary_pengelola=True,
        )

        for user in [active_user, pending_user]:
            refresh = RefreshToken.for_user(user)
            self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {refresh.access_token}")
            response = self.client.post(
                "/api/v1/invites/accept", {"token": invite_token}, format="json"
            )
            self.assertEqual(response.status_code, 400)
            self.assertEqual(response.data["error"], "Akun ini sudah tergabung dengan bank sampah")
            user.refresh_from_db()
            self.assertNotEqual(user.bank_sampah_id, target_bank.id)

    def test_invite_join_reassigns_rejected_primary_to_active_bank(self) -> None:
        target_bank = BankSampah.objects.create(
            nama="Target Join", alamat="Depok", no_hp_pic="+628111111111"
        )
        invite_token = target_bank.invite_token = "target-reassign-token"
        target_bank.invite_token_expires = timezone.now() + timedelta(days=3)
        target_bank.save(update_fields=["invite_token", "invite_token_expires", "updated_at"])
        rejected_bank = BankSampah.objects.create(
            nama="Rejected History",
            alamat="Depok",
            no_hp_pic="+628222222222",
            status=BankSampah.Status.REJECTED,
            is_active=False,
        )
        user = User.objects.create_user(
            email="rejected-primary@example.com",
            nama="Rejected Primary",
            no_hp="+628555555555",
            jenis_kelamin=User.Gender.MALE,
            tanggal_lahir="1990-01-01",
            bank_sampah=rejected_bank,
            is_profile_complete=True,
            is_primary_pengelola=True,
        )
        refresh = RefreshToken.for_user(user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {refresh.access_token}")

        response = self.client.post(
            "/api/v1/invites/accept", {"token": invite_token}, format="json"
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["outcome"], "join_success")
        self.assertEqual(response.data["message"], "Berhasil bergabung ke Bank Sampah Target Join")
        self.assertEqual(response.data["bank_sampah_id"], str(target_bank.id))
        self.assertEqual(response.data["bank_sampah_nama"], "Target Join")
        user.refresh_from_db()
        self.assertEqual(user.bank_sampah_id, target_bank.id)
        self.assertFalse(user.is_primary_pengelola)
        self.assertTrue(
            BankSampah.objects.filter(
                id=rejected_bank.id, status=BankSampah.Status.REJECTED
            ).exists()
        )

    @override_settings(PILAH_SUPERADMIN_EMAILS=("root@example.com",))
    def test_superadmin_cannot_use_pengelola_endpoints(self) -> None:
        self.client.credentials()
        response = self.client.post(
            "/api/v1/auth/google",
            {"id_token": "dev-superadmin:root@example.com:Root Admin"},
            format="json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["user"]["role"], "superadmin")
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {response.data['access_token']}")

        dashboard = self.client.get("/api/v1/dashboard/stats")
        self.assertEqual(dashboard.status_code, 403)

        approvals = self.client.get("/api/v1/superadmin/bank-sampah")
        self.assertEqual(approvals.status_code, 200)

    @override_settings(PILAH_SUPERADMIN_EMAILS=("queue-admin@example.com",))
    def test_superadmin_bank_queue_sorts_oldest_first(self) -> None:
        self.client.credentials()
        superadmin = User.objects.create_user(
            email="queue-admin@example.com",
            nama="Queue Admin",
            role=User.Role.SUPERADMIN,
            is_profile_complete=True,
        )
        refresh = RefreshToken.for_user(superadmin)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {refresh.access_token}")
        older = BankSampah.objects.create(
            nama="Older Pending",
            alamat="Depok",
            no_hp_pic="+628111111111",
            status=BankSampah.Status.PENDING,
        )
        newer = BankSampah.objects.create(
            nama="Newer Pending",
            alamat="Depok",
            no_hp_pic="+628222222222",
            status=BankSampah.Status.PENDING,
        )
        BankSampah.objects.filter(id=older.id).update(created_at=timezone.now() - timedelta(days=2))
        BankSampah.objects.filter(id=newer.id).update(created_at=timezone.now() - timedelta(days=1))

        response = self.client.get("/api/v1/superadmin/bank-sampah?status=pending")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["results"][0]["id"], str(older.id))
        self.assertEqual(response.data["results"][1]["id"], str(newer.id))

    def test_android_assetlinks(self) -> None:
        self.client.credentials()
        response = self.client.get("/.well-known/assetlinks.json")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/json")
        data = json.loads(response.content)
        self.assertEqual(
            data,
            [
                {
                    "relation": ["delegate_permission/common.handle_all_urls"],
                    "target": {
                        "namespace": "android_app",
                        "package_name": "com.mobile.pilahapp",
                        "sha256_cert_fingerprints": [
                            "B5:1A:6B:E6:CC:8F:00:0A:1E:BD:82:B9:6E:EB:80:68:B2:14:4D:FA:25:B5:B6:8A:E5:09:7D:DC:88:0F:70:36",
                            "73:B6:CC:52:38:29:88:E0:38:DA:47:1E:68:F9:86:6F:C8:C5:3E:8E:5D:1B:51:BD:FF:FE:78:ED:C6:C4:97:D6",
                            "A4:3F:BA:76:7F:D2:CA:A6:A9:8F:4A:99:66:78:47:AD:51:FB:97:6B:EE:C1:08:09:AC:EA:A8:64:A6:7A:36:3B",
                            "70:EF:3E:65:35:DD:83:3C:5B:43:79:E9:23:13:84:8E:E9:82:30:4F:A8:C6:E6:5F:A1:E5:3A:8A:6F:D8:EB:FC",
                        ],
                    },
                },
            ],
        )

    def test_pengelola_induk_can_sign_in(self) -> None:
        response = self.client.post(
            "/api/v1/auth/google",
            {"id_token": "dev-pengelola-induk:induk@example.com:Pengelola Induk"},
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["user"]["role"], "pengelola_induk")

    def test_nasabah_can_sign_in(self) -> None:
        response = self.client.post(
            "/api/v1/auth/google",
            {"id_token": "dev-nasabah:nasabah@example.com:Nasabah PILAH"},
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["user"]["role"], "nasabah")

    def test_new_role_logins_return_role_specific_states(self) -> None:
        expected_states = {
            "dev-pengelola-induk": "complete_profile",
            "dev-nasabah": "complete_profile",
        }
        for index, (token_prefix, state) in enumerate(expected_states.items(), start=1):
            with self.subTest(role=token_prefix):
                response = self.client.post(
                    "/api/v1/auth/google",
                    {"id_token": f"{token_prefix}:state{index}@example.com:State User"},
                    format="json",
                )

                self.assertEqual(response.status_code, 200, response.data)
                self.assertEqual(response.data["next_step"], state)
                self.assertEqual(response.data["user"]["state"], state)

    @override_settings(PILAH_ALLOW_FAKE_GOOGLE_TOKEN=True)
    def test_google_login_uses_role_assigned_to_existing_account(self) -> None:
        for role in (User.Role.PENGELOLA_INDUK, User.Role.NASABAH):
            with self.subTest(role=role):
                email = f"{role}@example.com"
                User.objects.create_user(email=email, nama=role, role=role)

                response = self.client.post(
                    "/api/v1/auth/google",
                    {"id_token": f"dev:{email}:Test User"},
                    format="json",
                )

                self.assertEqual(response.status_code, 200, response.data)
                self.assertEqual(response.data["user"]["role"], role)
                self.assertEqual(AccessToken(response.data["access_token"])["role"], role)
                self.assertEqual(response.data["next_step"], "complete_profile")

    def test_new_roles_do_not_inherit_pengelola_endpoint_access(self) -> None:
        for role in (User.Role.PENGELOLA_INDUK, User.Role.NASABAH):
            with self.subTest(role=role):
                user = User.objects.create_user(
                    email=f"{role}@example.com",
                    nama=role,
                    role=role,
                    bank_sampah=self.bank,
                    is_profile_complete=True,
                )
                refresh = RefreshToken.for_user(user)
                self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {refresh.access_token}")

                response = self.client.get("/api/v1/dashboard/stats")

                self.assertEqual(response.status_code, 403)

    @override_settings(PILAH_ALLOW_FAKE_GOOGLE_TOKEN=False)
    @patch("api.services.google_id_token.verify_oauth2_token")
    def test_google_profile_claim_cannot_assign_pilah_role(self, verify: Mock) -> None:
        verify.return_value = {
            "sub": "google-123",
            "email": "claim@example.com",
            "email_verified": True,
            "name": "Claim User",
            "role": User.Role.SUPERADMIN,
        }

        response = self.client.post(
            "/api/v1/auth/google", {"id_token": "signed-google-token"}, format="json"
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(response.data["registration_required"])
        self.assertFalse(User.objects.filter(email="claim@example.com").exists())

    @override_settings(PILAH_ALLOW_FAKE_GOOGLE_TOKEN=False)
    @patch("api.services.google_id_token.verify_oauth2_token")
    def test_google_profile_without_email_is_rejected(self, verify: Mock) -> None:
        verify.return_value = {"sub": "google-without-email", "name": "Missing Email"}

        response = self.client.post(
            "/api/v1/auth/google", {"id_token": "signed-google-token"}, format="json"
        )

        self.assertEqual(response.status_code, 401, response.data)
        self.assertEqual(response.data["error"], "ID Token invalid atau expired")

    def test_bank_induk_can_have_multiple_unit_banks(self) -> None:
        parent = BankSampah.objects.create(
            nama="Bank Sampah Induk",
            alamat="Depok",
            no_hp_pic="+628111111111",
            jenis_organisasi=BankSampah.OrganizationType.INDUK,
        )
        for number in (1, 2):
            BankSampah.objects.create(
                nama=f"Unit {number}",
                alamat="Depok",
                no_hp_pic=f"+62822222222{number}",
                jenis_organisasi=BankSampah.OrganizationType.UNIT,
                parent=parent,
            )

        self.assertEqual(parent.units.count(), 2)

    def test_mandiri_bank_can_be_validated_without_a_parent(self) -> None:
        bank = BankSampah(
            nama="Bank Mandiri",
            alamat="Depok",
            no_hp_pic="+628111111112",
        )

        bank.full_clean()

    def test_bank_unit_requires_an_induk_parent(self) -> None:
        unit = BankSampah(
            nama="Unit without induk",
            alamat="Depok",
            no_hp_pic="+628222222221",
            jenis_organisasi=BankSampah.OrganizationType.UNIT,
            parent=self.bank,
        )

        with self.assertRaises(ValidationError):
            unit.full_clean()

    def test_induk_with_units_cannot_be_demoted(self) -> None:
        parent = BankSampah.objects.create(
            nama="Bank Sampah Induk",
            alamat="Depok",
            no_hp_pic="+628111111111",
            jenis_organisasi=BankSampah.OrganizationType.INDUK,
        )
        BankSampah.objects.create(
            nama="Unit 1",
            alamat="Depok",
            no_hp_pic="+628222222221",
            jenis_organisasi=BankSampah.OrganizationType.UNIT,
            parent=parent,
        )

        parent.jenis_organisasi = BankSampah.OrganizationType.MANDIRI
        with self.assertRaises(ValidationError):
            parent.save()

    def test_unit_with_non_induk_parent_cannot_be_saved(self) -> None:
        with self.assertRaises(ValidationError):
            BankSampah.objects.create(
                nama="Invalid Unit",
                alamat="Depok",
                no_hp_pic="+628222222222",
                jenis_organisasi=BankSampah.OrganizationType.UNIT,
                parent=self.bank,
            )

    def test_bank_unit_cannot_be_saved_without_a_parent(self) -> None:
        with self.assertRaises(ValidationError):
            BankSampah.objects.create(
                nama="Unit without parent",
                alamat="Depok",
                no_hp_pic="+628222222222",
                jenis_organisasi=BankSampah.OrganizationType.UNIT,
            )

    def test_nasabah_account_can_join_multiple_banks(self) -> None:
        customer = User.objects.create_user(
            email="multi-bank-nasabah@example.com",
            nama="Nasabah PILAH",
            role=User.Role.NASABAH,
        )
        parent = BankSampah.objects.create(
            nama="Bank Sampah Induk Membership",
            alamat="Depok",
            no_hp_pic="+628111111113",
            jenis_organisasi=BankSampah.OrganizationType.INDUK,
        )
        banks = [
            BankSampah.objects.create(
                nama=f"Unit {number}",
                alamat="Depok",
                no_hp_pic=f"+62822222222{number}",
                jenis_organisasi=BankSampah.OrganizationType.UNIT,
                parent=parent,
            )
            for number in (1, 2)
        ]
        for number, bank in enumerate(banks, start=1):
            Nasabah.objects.create(
                user=customer,
                bank_sampah=bank,
                nomor=f"NAS-{number:04d}",
                nama=customer.nama,
                alamat="Depok",
                no_hp=f"+62833333333{number}",
            )

        self.assertEqual(customer.keanggotaan_nasabah.count(), 2)

    def test_nasabah_account_can_only_join_a_bank_once(self) -> None:
        customer = User.objects.create_user(
            email="single-bank-nasabah@example.com",
            nama="Nasabah PILAH",
            role=User.Role.NASABAH,
        )
        Nasabah.objects.create(
            user=customer,
            bank_sampah=self.bank,
            nomor="NAS-0001",
            nama=customer.nama,
            alamat="Depok",
            no_hp="+628333333331",
        )

        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Nasabah.objects.create(
                    user=customer,
                    bank_sampah=self.bank,
                    nomor="NAS-0002",
                    nama=customer.nama,
                    alamat="Depok",
                    no_hp="+628333333332",
                )

    def test_nasabah_membership_requires_a_nasabah_user(self) -> None:
        pengelola = User.objects.create_user(
            email="membership-pengelola@example.com",
            nama="Pengelola",
            role=User.Role.PENGELOLA,
        )

        with self.assertRaises(ValidationError):
            Nasabah.objects.create(
                user=pengelola,
                bank_sampah=self.bank,
                nomor="NAS-0003",
                nama=pengelola.nama,
                alamat="Depok",
                no_hp="+628333333333",
            )

    def test_nasabah_with_memberships_cannot_change_role(self) -> None:
        customer = User.objects.create_user(
            email="role-change-nasabah@example.com",
            nama="Nasabah PILAH",
            role=User.Role.NASABAH,
        )
        Nasabah.objects.create(
            user=customer,
            bank_sampah=self.bank,
            nomor="NAS-0004",
            nama=customer.nama,
            alamat="Depok",
            no_hp="+628333333334",
        )

        customer.role = User.Role.PENGELOLA
        with self.assertRaises(ValidationError) as raised:
            customer.save()

        self.assertEqual(
            raised.exception.message_dict["role"],
            ["Pengguna Nasabah yang masih memiliki keanggotaan tidak dapat berganti peran."],
        )
        customer.refresh_from_db()
        self.assertEqual(customer.role, User.Role.NASABAH)


class HealthzTests(TestCase):
    def test_healthz_ok(self) -> None:
        response = self.client.get("/healthz")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "ok")
        self.assertTrue(response.json()["database"])


class MediaRouteTests(TestCase):
    def test_local_media_route_uses_django_file_server(self) -> None:
        if settings.SERVE_MEDIA:
            match = resolve("/media/bank_sampah/logo/example.png")

            self.assertIs(match.func, serve)
            self.assertEqual(match.kwargs["path"], "example.png")
        else:
            with self.assertRaises(Resolver404):
                resolve("/media/bank_sampah/logo/example.png")

        with self.assertRaises(Resolver404):
            resolve("/media/bank_sampah/kegiatan/example.png")


class ProtectedMediaTests(TestCase):
    @override_settings(GS_BUCKET_NAME="")
    def test_activity_proof_uses_a_signed_url(self) -> None:
        bank = BankSampah.objects.create(
            nama="Test Bank",
            no_hp_pic="+628123456789",
            foto_kegiatan="bank_sampah/kegiatan/proof.png",
        )
        request = self.client.get("/").wsgi_request
        data = BankSampahApprovalListSerializer(bank, context={"request": request}).data

        self.assertIn("/media/activity/", data["foto_kegiatan"])
        token = data["foto_kegiatan"].split("/media/activity/", 1)[1]
        self.assertEqual(
            TimestampSigner(salt="bank-sampah-kegiatan").unsign(
                token, max_age=settings.MEDIA_SIGNED_URL_MAX_AGE
            ),
            "bank_sampah/kegiatan/proof.png",
        )

    @patch("api.views.default_storage")
    def test_signed_activity_proof_is_served(self, storage: Mock) -> None:
        storage.open.return_value = BytesIO(b"proof")
        token = TimestampSigner(salt="bank-sampah-kegiatan").sign("bank_sampah/kegiatan/proof.png")

        response = self.client.get(f"/media/activity/{token}")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(b"".join(cast(Any, response).streaming_content), b"proof")
        storage.open.assert_called_once_with("bank_sampah/kegiatan/proof.png", "rb")

    def test_unsigned_activity_proof_is_rejected(self) -> None:
        response = self.client.get("/media/activity/not-a-valid-token")

        self.assertEqual(response.status_code, 404)
