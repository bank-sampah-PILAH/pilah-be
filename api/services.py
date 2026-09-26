import secrets
from calendar import monthrange
from collections.abc import Mapping
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any, cast
from uuid import UUID

from django.conf import settings
from django.db import transaction
from django.db.models import QuerySet, Sum
from django.db.models.functions import Coalesce
from django.http import HttpRequest
from django.utils import timezone
from google.auth.transport import requests as google_requests
from google.oauth2 import id_token as google_id_token
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
from rest_framework import serializers
from rest_framework_simplejwt.tokens import RefreshToken

from api.models import (
    BankSampah,
    BankSampahApprovalLog,
    DetailTransaksi,
    JenisSampah,
    Nasabah,
    NasabahApprovalLog,
    Saldo,
    Transaksi,
    User,
)

# ponytail: compat shims — canonical homes are apps.notify.services and
# shared_kernel.numbering.
from apps.notify.services import (  # noqa: F401
    DEFAULT_WA_TEMPLATE,
    WhatsAppService,
)
from shared_kernel.numbering import NumberingService  # noqa: F401


class AuthService:
    @staticmethod
    def login_with_google(raw_id_token: str) -> dict[str, Any]:
        profile = AuthService._verify_google_token(raw_id_token)
        email = profile["email"]
        google_id = profile["sub"]
        name = profile.get("name") or email.split("@")[0]
        # Google proves identity, not a user's PILAH authorization.
        role = profile.get("_pilah_dev_role", User.Role.PENGELOLA)

        user = User.objects.filter(email=email).first()
        is_new_user = user is None
        if user is None:
            user = User.objects.create_user(
                email=email,
                google_id=google_id,
                nama=name,
                role=role,
                is_profile_complete=False,
            )
        else:
            updates = []
            nama_was_empty = not user.nama
            if not user.google_id:
                user.google_id = google_id
                updates.append("google_id")
            if not user.nama:
                user.nama = name
                updates.append("nama")
            if updates:
                user.save(update_fields=updates)

        # PIL-154: nasabah added by pengurus via email syncs with the Google
        # account on first login — prefill empty User fields, never overwrite.
        # The token-derived nama yields to the pengurus-entered nasabah name
        # whenever it was empty before this sync (new accounts included).
        nasabah = Nasabah.objects.filter(email__iexact=user.email, is_active=True).first()
        if nasabah:
            profile_updates = []
            for user_field, nasabah_field in (
                ("nama", "nama"),
                ("no_hp", "no_hp"),
                ("jenis_kelamin", "jenis_kelamin"),
                ("tanggal_lahir", "tanggal_lahir"),
            ):
                fillable = user_field == "nama" and (is_new_user or nama_was_empty)
                if (fillable or not getattr(user, user_field)) and getattr(nasabah, nasabah_field):
                    setattr(user, user_field, getattr(nasabah, nasabah_field))
                    profile_updates.append(user_field)
            # Only a fully-populated profile counts as complete — a nasabah
            # record may itself be missing jenis_kelamin/tanggal_lahir, and
            # flagging complete here would lock /onboarding/profile out.
            if profile_updates and all(
                getattr(user, f) for f in ("nama", "no_hp", "jenis_kelamin", "tanggal_lahir")
            ):
                user.is_profile_complete = True
                profile_updates.append("is_profile_complete")
            if profile_updates:
                user.save(update_fields=profile_updates)

        refresh = RefreshToken.for_user(user)
        access = refresh.access_token
        if user.bank_sampah_id:
            access["bank_sampah_id"] = str(user.bank_sampah_id)
        access["role"] = user.role
        access["email"] = user.email
        bank = user.bank_sampah
        state = AuthService.user_state(user)
        return {
            "access_token": str(access),
            "refresh_token": str(refresh),
            "token_type": "Bearer",
            "expires_in": int(
                cast(timedelta, settings.SIMPLE_JWT["ACCESS_TOKEN_LIFETIME"]).total_seconds()
            ),
            "user": {
                "id": str(user.id),
                "name": user.nama,
                "nama": user.nama,
                "email": user.email,
                "role": user.role,
                "bank_sampah_id": str(user.bank_sampah_id) if user.bank_sampah_id else None,
                "bank_sampah_nama": bank.nama if bank else None,
                "bank_sampah_status": bank.status if bank else None,
                "is_profile_complete": user.is_profile_complete,
                "is_primary_pengelola": user.is_primary_pengelola,
                "state": state,
            },
            "next_step": state,
            "is_new_user": is_new_user,
        }

    @staticmethod
    def user_state(user: User) -> str:
        if user.role == User.Role.SUPERADMIN:
            return "superadmin_dashboard"
        if user.role == User.Role.PENGELOLA_INDUK:
            return "pengelola_induk_dashboard"
        if user.role == User.Role.NASABAH:
            return "nasabah_dashboard"
        if not user.is_profile_complete:
            return "complete_profile"
        bank = user.bank_sampah
        if not user.bank_sampah_id or bank is None:
            return "register_bank_sampah"
        if bank.status == BankSampah.Status.PENDING:
            return "approval_pending"
        if bank.status == BankSampah.Status.REJECTED:
            return "registration_rejected"
        return "dashboard"

    @staticmethod
    def _verify_google_token(raw_id_token: str) -> Mapping[str, Any]:
        if settings.PILAH_ALLOW_FAKE_GOOGLE_TOKEN:
            dev_roles = {
                "dev-superadmin:": ("dev-superadmin", User.Role.SUPERADMIN),
                "dev-pengelola-induk:": ("dev-pengelola-induk", User.Role.PENGELOLA_INDUK),
                "dev-nasabah:": ("dev-nasabah", User.Role.NASABAH),
            }
            for prefix, (subject_prefix, role) in dev_roles.items():
                if raw_id_token.startswith(prefix):
                    _, email, name = (raw_id_token.split(":", 2) + [""])[:3]
                    return {
                        "sub": f"{subject_prefix}-{email}",
                        "email": email,
                        "name": name or email.split("@")[0],
                        "_pilah_dev_role": role,
                    }
        if settings.PILAH_ALLOW_FAKE_GOOGLE_TOKEN and raw_id_token.startswith("dev:"):
            _, email, name = (raw_id_token.split(":", 2) + [""])[:3]
            return {"sub": f"dev-{email}", "email": email, "name": name or email.split("@")[0]}
        audience = settings.GOOGLE_CLIENT_ID or None
        try:
            profile = google_id_token.verify_oauth2_token(  # type: ignore[no-untyped-call]  # google-auth ships no stubs
                raw_id_token, google_requests.Request(), audience
            )
        except Exception as exc:
            raise serializers.ValidationError(
                {"id_token": ["ID Token invalid atau expired"]}
            ) from exc
        verified_profile = cast(Mapping[str, Any], profile)
        profile_subject = verified_profile.get("sub")
        profile_email = verified_profile.get("email")
        if (
            not isinstance(profile_subject, str)
            or not profile_subject
            or not isinstance(profile_email, str)
            or not profile_email
        ):
            raise serializers.ValidationError(
                {"id_token": ["ID Token tidak memuat email atau subject yang diperlukan"]}
            )
        return {
            "sub": profile_subject,
            "email": profile_email,
            "name": verified_profile.get("name"),
        }


class OnboardingService:
    @staticmethod
    @transaction.atomic
    def complete_profile(user: User, profile_data: Mapping[str, Any]) -> User:
        if user.is_profile_complete:
            raise ValueError("Profil sudah lengkap")
        for field, value in profile_data.items():
            setattr(user, field, value)
        user.is_profile_complete = True
        user.save(
            update_fields=[
                "nama",
                "no_hp",
                "jenis_kelamin",
                "tanggal_lahir",
                "is_profile_complete",
                "updated_at",
            ]
        )
        return user

    @staticmethod
    @transaction.atomic
    def register_bank_sampah(user: User, payload: Mapping[str, Any]) -> BankSampah:
        if user.role != User.Role.PENGELOLA:
            raise PermissionError("Hanya pengelola yang dapat mendaftarkan bank sampah")
        bank = user.bank_sampah
        if bank and bank.status == BankSampah.Status.ACTIVE:
            raise ValueError("Bank sampah sudah aktif")
        if bank and bank.status == BankSampah.Status.PENDING:
            raise ValueError("Pendaftaran bank sampah sedang diproses")
        if bank is None:
            bank = BankSampah()
        bank.nama = payload["nama"]
        bank.alamat = payload["alamat"]
        bank.kota = payload.get("kota", "")
        bank.no_hp_pic = payload["no_hp_pic"]
        bank.foto_kegiatan = payload["foto_kegiatan"]
        bank.status = BankSampah.Status.PENDING
        bank.is_active = False
        bank.wa_template = bank.wa_template or DEFAULT_WA_TEMPLATE
        bank.save()
        user.bank_sampah = bank
        user.is_primary_pengelola = True
        user.save(update_fields=["bank_sampah", "is_primary_pengelola", "updated_at"])
        return bank

    @staticmethod
    @transaction.atomic
    def accept_invite(user: User, token: str) -> tuple[BankSampah, str]:
        bank = BankSampah.objects.filter(
            invite_token=token, invite_token_expires__gt=timezone.now()
        ).first()
        if not bank or bank.status != BankSampah.Status.ACTIVE:
            raise ValueError("Tautan undangan tidak valid atau sudah kedaluwarsa")
        if user.role != User.Role.PENGELOLA:
            raise PermissionError("Hanya pengelola yang dapat menerima undangan")
        if user.bank_sampah_id == bank.id:
            return bank, "already_member"
        user_bank = user.bank_sampah
        if user_bank and user_bank.status in [
            BankSampah.Status.ACTIVE.value,
            BankSampah.Status.PENDING.value,
        ]:
            raise ValueError("Akun ini sudah tergabung dengan bank sampah")
        outcome = "join_success"
        user.bank_sampah = bank
        user.is_primary_pengelola = False
        if user.nama and user.no_hp and user.jenis_kelamin and user.tanggal_lahir:
            user.is_profile_complete = True
        user.save(
            update_fields=[
                "bank_sampah",
                "is_primary_pengelola",
                "is_profile_complete",
                "updated_at",
            ]
        )
        return bank, outcome


class ApprovalService:
    @staticmethod
    @transaction.atomic
    def approve(bank: BankSampah, superadmin: User, catatan: str = "") -> BankSampahApprovalLog:
        bank.status = BankSampah.Status.ACTIVE
        bank.is_active = True
        bank.save(update_fields=["status", "is_active", "updated_at"])
        return BankSampahApprovalLog.objects.create(
            bank_sampah=bank,
            superadmin=superadmin,
            status=BankSampahApprovalLog.Status.APPROVED,
            catatan=catatan,
        )

    @staticmethod
    @transaction.atomic
    def reject(bank: BankSampah, superadmin: User, catatan: str = "") -> BankSampahApprovalLog:
        bank.status = BankSampah.Status.REJECTED
        bank.is_active = False
        bank.save(update_fields=["status", "is_active", "updated_at"])
        return BankSampahApprovalLog.objects.create(
            bank_sampah=bank,
            superadmin=superadmin,
            status=BankSampahApprovalLog.Status.REJECTED,
            catatan=catatan,
        )


class TeamService:
    @staticmethod
    def generate_invite(bank: BankSampah) -> str:
        bank.invite_token = secrets.token_urlsafe(32)
        bank.invite_token_expires = timezone.now() + timedelta(days=3)
        bank.save(update_fields=["invite_token", "invite_token_expires", "updated_at"])
        return bank.invite_token


class NasabahApprovalService:
    @staticmethod
    @transaction.atomic
    def approve(nasabah: Nasabah, pengurus: User, catatan: str = "") -> NasabahApprovalLog:
        nasabah.status = Nasabah.Status.APPROVED
        nasabah.is_active = True
        nasabah.save(update_fields=["status", "is_active", "updated_at"])
        return NasabahApprovalLog.objects.create(
            nasabah=nasabah,
            pengurus=pengurus,
            status=NasabahApprovalLog.Status.APPROVED,
            catatan=catatan,
        )

    @staticmethod
    @transaction.atomic
    def reject(nasabah: Nasabah, pengurus: User, catatan: str = "") -> NasabahApprovalLog:
        nasabah.status = Nasabah.Status.REJECTED
        nasabah.is_active = False
        nasabah.save(update_fields=["status", "is_active", "updated_at"])
        return NasabahApprovalLog.objects.create(
            nasabah=nasabah,
            pengurus=pengurus,
            status=NasabahApprovalLog.Status.REJECTED,
            catatan=catatan,
        )


class TransactionService:
    @staticmethod
    @transaction.atomic
    def create_setoran(user: User, payload: Mapping[str, Any]) -> Transaksi:
        bank = user.bank_sampah
        assert bank is not None  # ponytail: views gate on IsActivePengelola
        nasabah = (
            Nasabah.objects.select_for_update()
            .filter(
                id=payload["nasabah_id"],
                bank_sampah=bank,
                is_active=True,
                status=Nasabah.Status.APPROVED,
            )
            .first()
        )
        if not nasabah:
            raise serializers.ValidationError(
                {"nasabah_id": ["Nasabah tidak ditemukan atau tidak aktif"]}
            )

        saldo, _ = Saldo.objects.select_for_update().get_or_create(nasabah=nasabah)
        transaksi = Transaksi.objects.create(
            nasabah=nasabah,
            bank_sampah=bank,
            dicatat_oleh=user,
            catatan=payload.get("catatan") or None,
        )

        total_nilai = Decimal("0.00")
        for index, item_payload in enumerate(payload["items"]):
            jenis = JenisSampah.objects.filter(
                id=item_payload["jenis_sampah_id"], bank_sampah=bank, is_active=True
            ).first()
            if not jenis:
                raise serializers.ValidationError(
                    {
                        f"items[{index}].jenis_sampah_id": [
                            "Jenis sampah tidak ditemukan atau tidak aktif"
                        ]
                    }
                )
            harga = item_payload.get("harga_per_kg") or jenis.harga_per_kg
            berat = item_payload["berat"]
            subtotal = (harga * berat).quantize(Decimal("0.01"))
            DetailTransaksi.objects.create(
                transaksi=transaksi,
                jenis_sampah=jenis,
                nama_sampah_snapshot=jenis.nama_sampah,
                kategori_snapshot=jenis.kategori,
                harga_snapshot=harga,
                berat=berat,
                subtotal=subtotal,
            )
            total_nilai += subtotal

        transaksi.total_nilai = total_nilai
        transaksi.save(update_fields=["total_nilai"])
        saldo.total_saldo += total_nilai
        saldo.save(update_fields=["total_saldo", "updated_at"])
        return transaksi

    @staticmethod
    def export_excel(
        queryset: QuerySet[Transaksi], request: HttpRequest | None = None
    ) -> tuple[bytes, str]:
        return _export_excel(queryset, request)


class TransactionFilterService:
    @staticmethod
    def apply_period(queryset: QuerySet[Transaksi], request: HttpRequest) -> QuerySet[Transaksi]:
        periode = request.GET.get("periode", "hari_ini")
        today = timezone.localdate()
        if periode == "hari_ini":
            start, end = today, today
        elif periode == "minggu_ini":
            start, end = today - timedelta(days=today.weekday()), today
        elif periode == "bulan_ini":
            start, end = today.replace(day=1), today
        elif periode == "bulan_lalu":
            first_this_month = today.replace(day=1)
            last_previous_month = first_this_month - timedelta(days=1)
            start = last_previous_month.replace(day=1)
            end = last_previous_month
        elif periode == "custom":
            parsed_start = TransactionFilterService._parse_date(request.GET.get("dari_tanggal"))
            parsed_end = TransactionFilterService._parse_date(request.GET.get("sampai_tanggal"))
            if not parsed_start or not parsed_end:
                raise serializers.ValidationError(
                    {"error": "dari_tanggal dan sampai_tanggal wajib diisi"}
                )
            if parsed_end < parsed_start:
                raise ValueError("Tanggal akhir tidak boleh lebih awal dari tanggal awal")
            start, end = parsed_start, parsed_end
        else:
            return queryset

        tz = timezone.get_current_timezone()
        start_dt = timezone.make_aware(datetime.combine(start, time.min), tz)
        end_dt = timezone.make_aware(datetime.combine(end, time.max), tz)
        return queryset.filter(tanggal__range=(start_dt, end_dt))

    @staticmethod
    def _parse_date(value: str | None) -> date | None:
        if not value:
            return None
        try:
            return date.fromisoformat(value)
        except ValueError as exc:
            raise serializers.ValidationError({"error": "Format tanggal harus YYYY-MM-DD"}) from exc


class DashboardService:
    @staticmethod
    def stats(user: User) -> dict[str, Any]:
        bank = user.bank_sampah
        assert bank is not None  # ponytail: views gate on IsActivePengelola
        today = timezone.localdate()
        start = today.replace(day=1)
        end = today.replace(day=monthrange(today.year, today.month)[1])
        tz = timezone.get_current_timezone()
        start_dt = timezone.make_aware(datetime.combine(start, time.min), tz)
        end_dt = timezone.make_aware(datetime.combine(end, time.max), tz)

        transaksi = Transaksi.objects.filter(bank_sampah=bank, tanggal__range=(start_dt, end_dt))
        totals = DetailTransaksi.objects.filter(transaksi__in=transaksi).aggregate(
            total_kg=Coalesce(Sum("berat"), Decimal(0)),
        )
        nilai = transaksi.aggregate(total=Coalesce(Sum("total_nilai"), Decimal(0)))["total"]
        return {
            "bank_sampah_nama": bank.nama,
            "pengelola_nama": user.nama,
            "periode": today.strftime("%Y-%m"),
            "nasabah_aktif": Nasabah.objects.filter(
                bank_sampah=bank, is_active=True, status=Nasabah.Status.APPROVED
            ).count(),
            "transaksi_bulan_ini": transaksi.count(),
            "total_sampah_kg_bulan_ini": totals["total_kg"],
            "total_nilai_bulan_ini": nilai,
        }


def _month_label() -> str:
    return timezone.localdate().strftime("%B_%Y")


class TransactionExportMixin:
    pass


def _export_excel(
    queryset: QuerySet[Transaksi], request: HttpRequest | None = None
) -> tuple[bytes, str]:
    wb = Workbook()
    _fill_summary_export_sheet(cast(Worksheet, wb.active), queryset)
    _fill_raw_export_sheet(wb.create_sheet("Riwayat Transaksi"), queryset, request)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue(), f"PILAH_Laporan_{_month_label()}.xlsx"


def _fill_summary_export_sheet(ws: Worksheet, queryset: QuerySet[Transaksi]) -> None:
    ws.title = "Laporan"
    ws.append(["Jenis", "Sampah", "Harga per kg", "Jumlah kg", "Total"])
    rows = (
        DetailTransaksi.objects.filter(transaksi__in=queryset)
        .values("kategori_snapshot", "nama_sampah_snapshot", "harga_snapshot")
        .annotate(total_kg=Sum("berat"))
        .order_by("kategori_snapshot", "nama_sampah_snapshot")
    )
    current_row = 2
    for row in rows:
        ws.append(
            [
                row["kategori_snapshot"],
                row["nama_sampah_snapshot"],
                float(row["harga_snapshot"]),
                float(row["total_kg"]),
                f"=C{current_row}*D{current_row}",
            ]
        )
        current_row += 1
    ws.append(["", "", "", "TOTAL", f"=SUM(E2:E{current_row - 1})"])
    for column in ["A", "B", "C", "D", "E"]:
        ws.column_dimensions[column].width = 22


def _fill_raw_export_sheet(
    ws: Worksheet, queryset: QuerySet[Transaksi], request: HttpRequest | None = None
) -> None:
    ws.title = "Riwayat Transaksi"
    ws.merge_cells("A1:J1")
    ws.merge_cells("A2:J2")
    ws["A1"] = "PILAH - Riwayat Transaksi"
    ws["A2"] = _export_filter_label(queryset, request)
    ws.append([])
    ws.append(
        [
            "No",
            "Tanggal",
            "Waktu",
            "Nama Nasabah",
            "ID Nasabah",
            "Jenis Sampah",
            "Berat (kg)",
            "Harga/kg (Rp)",
            "Subtotal (Rp)",
            "Saldo Setelah Transaksi (Rp)",
        ]
    )

    details = (
        DetailTransaksi.objects.filter(transaksi__in=queryset)
        .select_related("transaksi__nasabah")
        .order_by("-transaksi__tanggal", "id")
    )
    saldo_after_by_transaction = _saldo_after_by_transaction(queryset)
    for number, item in enumerate(details, start=1):
        local_datetime = timezone.localtime(item.transaksi.tanggal)
        ws.append(
            [
                number,
                local_datetime.strftime("%d/%m/%Y"),
                local_datetime.strftime("%H:%M"),
                item.transaksi.nasabah.nama,
                item.transaksi.nasabah.nomor,
                item.nama_sampah_snapshot,
                float(item.berat),
                int(item.harga_snapshot),
                int(item.subtotal),
                int(saldo_after_by_transaction[item.transaksi_id]),
            ]
        )

    _style_raw_export_sheet(ws)


def _saldo_after_by_transaction(queryset: QuerySet[Transaksi]) -> dict[UUID, Decimal]:
    target_transactions = list(queryset.select_related("bank_sampah").order_by("tanggal", "id"))
    if not target_transactions:
        return {}

    target_ids = {trans.id for trans in target_transactions}
    nasabah_ids = {trans.nasabah_id for trans in target_transactions}
    latest_transaction = target_transactions[-1]
    bank_sampah = latest_transaction.bank_sampah
    running_balances = {nasabah_id: Decimal("0.00") for nasabah_id in nasabah_ids}
    saldo_after = {}

    transactions = (
        Transaksi.objects.filter(
            bank_sampah=bank_sampah,
            nasabah_id__in=nasabah_ids,
            tanggal__lte=latest_transaction.tanggal,
        )
        .only("id", "nasabah_id", "total_nilai", "tanggal")
        .order_by("nasabah_id", "tanggal", "id")
    )
    for trans in transactions:
        running_balances[trans.nasabah_id] += trans.total_nilai
        if trans.id in target_ids:
            saldo_after[trans.id] = running_balances[trans.nasabah_id]

    return saldo_after


def _export_filter_label(queryset: QuerySet[Transaksi], request: HttpRequest | None) -> str:
    first_transaction = queryset.first()
    bank_name = first_transaction.bank_sampah.nama if first_transaction else "-"
    downloaded_at = timezone.localtime().strftime("%d/%m/%Y %H:%M")
    period_label = "Semua Periode"
    if request:
        period_label = _period_label(request)
    return f"Filter periode: {period_label}   |   Bank Sampah: {bank_name}   |   Diunduh: {downloaded_at}"


def _period_label(request: HttpRequest) -> str:
    periode = request.GET.get("periode", "hari_ini")
    today = timezone.localdate()
    if periode == "hari_ini":
        return f"Hari Ini ({today.strftime('%d/%m/%Y')})"
    if periode == "minggu_ini":
        start = today - timedelta(days=today.weekday())
        return f"Minggu Ini ({start.strftime('%d/%m/%Y')} - {today.strftime('%d/%m/%Y')})"
    if periode == "bulan_ini":
        end_day = monthrange(today.year, today.month)[1]
        return f"Bulan Ini (1 - {end_day} {_month_name(today.month)} {today.year})"
    if periode == "bulan_lalu":
        first_this_month = today.replace(day=1)
        previous_month = first_this_month - timedelta(days=1)
        end_day = monthrange(previous_month.year, previous_month.month)[1]
        return (
            f"Bulan Lalu (1 - {end_day} {_month_name(previous_month.month)} {previous_month.year})"
        )
    if periode == "custom":
        start_date = request.GET.get("dari_tanggal", "")
        end_date = request.GET.get("sampai_tanggal", "")
        return f"Custom ({start_date} - {end_date})"
    return "Semua Periode"


def _month_name(month: int) -> str:
    return [
        "Januari",
        "Februari",
        "Maret",
        "April",
        "Mei",
        "Juni",
        "Juli",
        "Agustus",
        "September",
        "Oktober",
        "November",
        "Desember",
    ][month - 1]


def _style_raw_export_sheet(ws: Worksheet) -> None:
    ws.freeze_panes = "A5"
    widths = {
        "A": 5,
        "B": 13,
        "C": 9,
        "D": 18,
        "E": 13,
        "F": 16,
        "G": 11,
        "H": 13,
        "I": 13,
        "J": 28,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    title_fill = PatternFill("solid", fgColor="166534")
    subtitle_fill = PatternFill("solid", fgColor="DCFCE7")
    header_fill = PatternFill("solid", fgColor="16A34A")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A2"].font = Font(size=10, color="374151")
    ws["A2"].fill = subtitle_fill
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 22

    for row in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=10):
        for cell in row:
            cell.fill = title_fill if cell.row == 1 else subtitle_fill

    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=10):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
        row[0].alignment = Alignment(horizontal="center", vertical="center")
        row[6].number_format = "0.##"
        for cell in row[7:10]:
            cell.number_format = "#,##0"
