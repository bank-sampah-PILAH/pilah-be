import base64
import re
import zlib
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from api.models import DraftPencairan, User
from apps.ledger.draft_testing import URL, DraftTestBase

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _teks_pdf(content: bytes) -> str:
    """Text drawn on the PDF pages: reportlab Flate+ASCII85-encodes every content stream."""
    teks = []
    for aliran in re.findall(rb"stream\r?\n(.*?)endstream", content, re.DOTALL):
        try:
            teks.append(
                zlib.decompress(base64.a85decode(aliran.strip().removesuffix(b"~>"))).decode(
                    "latin-1"
                )
            )
        except (ValueError, zlib.error):
            continue
    return "\n".join(teks)


class ExportDraftTests(DraftTestBase):
    def setUp(self) -> None:
        super().setUp()
        self.budi = self._nasabah("NAS-0002", "Budi Santoso", "50000")
        response = self._buat(
            [
                {"nasabah_id": str(self.nasabah.id), "nominal": "100000", "metode": "transfer"},
                {
                    "nasabah_id": str(self.budi.id),
                    "potongan_jenis": "rupiah",
                    "potongan_nilai": "1500",
                },
            ],
            nama="Cair Oktober",
            potongan_jenis="persen",
            potongan_nilai="10",
        )
        self.draft = response.data

    def _export(self, berkas: str | None, draft_id: str | None = None) -> Any:
        params = {"berkas": berkas} if berkas else {}
        return self.client.get(f"{URL}/{draft_id or self.draft['id']}/export", params)

    def test_export_excel_memuat_rincian_per_nasabah_dan_total(self) -> None:
        response = self._export("xlsx")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX)
        self.assertIn("attachment; filename=", response["Content-Disposition"])
        self.assertIn(".xlsx", response["Content-Disposition"])
        baris = [
            [sel for sel in row if sel is not None]
            for row in load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True)
        ]
        self.assertIn("Cair Oktober", [sel for row in baris for sel in row])
        self.assertIn([1, "NAS-0001", "Ahmad Ridwan", "Transfer", 100000, 10000, 90000], baris)
        self.assertIn([2, "NAS-0002", "Budi Santoso", "Tunai", 50000, 1500, 48500], baris)
        self.assertIn(["TOTAL", 150000, 11500, 138500], baris)

    def test_export_pdf_memuat_nama_nasabah_dan_total(self) -> None:
        response = self._export("pdf")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertIn(".pdf", response["Content-Disposition"])
        teks = _teks_pdf(response.content)
        for isi in ("Cair Oktober", "Ahmad Ridwan", "Budi Santoso", "90.000", "48.500", "138.500"):
            self.assertIn(isi, teks)

    def test_export_tidak_mengubah_status_dan_bisa_diulang_setelah_konfirmasi(self) -> None:
        self.assertEqual(self._export("pdf").status_code, 200)
        self.assertEqual(self._export("xlsx").status_code, 200)
        self.assertEqual(DraftPencairan.objects.get().status, "draft")

        self.client.post(f"{URL}/{self.draft['id']}/konfirmasi")

        self.assertEqual(self._export("pdf").status_code, 200)
        self.assertEqual(self._export("xlsx").status_code, 200)
        self.assertEqual(DraftPencairan.objects.get().status, "dikonfirmasi")

    def test_berkas_wajib_dan_hanya_pdf_atau_xlsx(self) -> None:
        for berkas in (None, "csv", ""):
            with self.subTest(berkas=berkas):
                self.assertEqual(self._export(berkas).status_code, 422)

    def test_export_draft_bank_lain_atau_oleh_nasabah_ditolak(self) -> None:
        asing = self._draft_bank_lain()
        self.assertEqual(self._export("pdf", str(asing.id)).status_code, 404)
        self._login(
            User.objects.create_user(email="n@example.test", nama="N", role=User.Role.NASABAH)
        )

        self.assertEqual(self._export("pdf").status_code, 403)


class ExportPratinjauTests(DraftTestBase):
    """Export of what is on the screen, saved or not: nothing is written."""

    def setUp(self) -> None:
        super().setUp()
        self.budi = self._nasabah("NAS-0002", "Budi Santoso", "50000")

    def _body(self, **extra: Any) -> dict[str, Any]:
        return {
            "nama": "Belum Disimpan",
            "potongan_jenis": "persen",
            "potongan_nilai": "10",
            "items": [
                {"nasabah_id": str(self.nasabah.id), "nominal": "100000"},
                {"nasabah_id": str(self.budi.id)},
            ],
            **extra,
        }

    def _export(self, body: dict[str, Any], berkas: str | None = "pdf") -> Any:
        params = f"?berkas={berkas}" if berkas else ""
        return self.client.post(f"{URL}/export{params}", body, format="json")

    def test_pdf_dari_isi_yang_dikirim_tanpa_menyimpan_apa_pun(self) -> None:
        response = self._export(self._body())

        self.assertEqual(response.status_code, 200, getattr(response, "data", None))
        self.assertEqual(response["Content-Type"], "application/pdf")
        teks = _teks_pdf(response.content)
        for isi in ("Belum Disimpan", "Ahmad Ridwan", "Budi Santoso", "90.000", "45.000"):
            self.assertIn(isi, teks)
        self.assertEqual(DraftPencairan.objects.count(), 0)

    def test_excel_dari_isi_yang_dikirim(self) -> None:
        response = self._export(self._body(), "xlsx")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX)
        baris = [
            [sel for sel in row if sel is not None]
            for row in load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True)
        ]
        self.assertIn([1, "NAS-0001", "Ahmad Ridwan", "Tunai", 100000, 10000, 90000], baris)
        self.assertEqual(DraftPencairan.objects.count(), 0)

    def test_tidak_menyentuh_draft_yang_sudah_tersimpan(self) -> None:
        tersimpan = self._buat(
            [{"nasabah_id": str(self.nasabah.id), "nominal": "100000"}], nama="Versi Tersimpan"
        ).data

        self._export(self._body())

        draft = DraftPencairan.objects.get()
        self.assertEqual(str(draft.id), tersimpan["id"])
        self.assertEqual(draft.nama, "Versi Tersimpan")
        self.assertEqual(draft.items.count(), 1)

    def test_nama_kosong_dipakai_nama_default(self) -> None:
        response = self._export(self._body(nama=""))

        self.assertIn("Pencairan", _teks_pdf(response.content))

    def test_isi_tidak_valid_ditolak_seperti_menyimpan(self) -> None:
        body = self._body(items=[{"nasabah_id": str(self.budi.id), "nominal": "99999999"}])

        response = self._export(body)

        self.assertEqual(response.status_code, 422)
        self.assertEqual(DraftPencairan.objects.count(), 0)

    def test_berkas_wajib_dan_hanya_pdf_atau_xlsx(self) -> None:
        for berkas in (None, "csv"):
            with self.subTest(berkas=berkas):
                self.assertEqual(self._export(self._body(), berkas).status_code, 422)

    def test_nasabah_tidak_boleh_mengekspor(self) -> None:
        self._login(
            User.objects.create_user(email="n@example.test", nama="N", role=User.Role.NASABAH)
        )

        self.assertIn(self._export(self._body()).status_code, (403, 404))
