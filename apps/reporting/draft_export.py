"""PDF and Excel files for a draft pencairan (PIL-316).

Both formats carry the same rows: one per nasabah with nominal, effective potongan and the
amount to pay, then the batch total. Money is whole rupiah, so cells hold plain integers.
"""

from decimal import Decimal
from io import BytesIO
from typing import cast
from xml.sax.saxutils import escape

from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from api.models import DraftPencairan
from shared_kernel.kalkulasi import format_ribuan

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_TYPE = "application/pdf"

_Baris = tuple[int, str, str, str, int, int, int]

_HEADER = ["No", "Kode", "Nama", "Metode", "Nominal (Rp)", "Potongan (Rp)", "Dibayar (Rp)"]


def export_draft_pencairan(draft: DraftPencairan, berkas: str) -> tuple[bytes, str, str]:
    """Return (content, filename, content_type) for `berkas` in {"pdf", "xlsx"}."""
    stem = f"PILAH_Draft_Pencairan_{slugify(draft.nama) or 'draft'}_{timezone.localdate():%Y%m%d}"
    if berkas == "pdf":
        return _pdf(draft), f"{stem}.pdf", PDF_TYPE
    return _xlsx(draft), f"{stem}.xlsx", XLSX_TYPE


def _rows(draft: DraftPencairan) -> list[_Baris]:
    return [
        (
            no,
            item.nasabah.nomor,
            item.nasabah.nama,
            item.get_metode_display(),
            int(item.nominal),
            int(item.potongan),
            int(item.dibayar),
        )
        for no, item in enumerate(draft.items.select_related("nasabah", "draft"), start=1)
    ]


def _totals(draft: DraftPencairan) -> list[int]:
    return [int(draft.total_nominal), int(draft.total_potongan), int(draft.total_dibayar)]


def _meta(draft: DraftPencairan) -> list[tuple[str, str]]:
    dibuat = timezone.localtime(draft.created_at).strftime("%d/%m/%Y %H:%M")
    return [
        ("Nama draft", draft.nama),
        ("Status", draft.get_status_display()),
        ("Dibuat oleh", f"{draft.dibuat_oleh.nama} ({dibuat})"),
    ]


def _xlsx(draft: DraftPencairan) -> bytes:
    wb = Workbook()
    ws = cast(Worksheet, wb.active)
    ws.title = "Draft Pencairan"
    ws.append(["Draft Pencairan"])
    ws["A1"].font = Font(bold=True, size=14)
    for label, nilai in _meta(draft):
        ws.append([label, nilai])
    ws.append([])
    ws.append(_HEADER)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="006D44")
        cell.alignment = Alignment(horizontal="center")
    for row in _rows(draft):
        ws.append(list(row))
    ws.append(["TOTAL", None, None, None, *_totals(draft)])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for kolom, lebar in zip("ABCDEFG", (14, 12, 28, 12, 16, 16, 16), strict=True):
        ws.column_dimensions[kolom].width = lebar
    for baris in ws.iter_rows(min_col=5, max_col=7):
        for cell in baris:
            cell.number_format = "#,##0"
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _pdf(draft: DraftPencairan) -> bytes:
    gaya = getSampleStyleSheet()
    isi: list[Paragraph | Spacer | Table] = [
        Paragraph("Draft Pencairan", gaya["Title"]),
        *(
            Paragraph(f"<b>{escape(label)}:</b> {escape(nilai)}", gaya["Normal"])
            for label, nilai in _meta(draft)
        ),
        Spacer(1, 0.5 * cm),
    ]
    tabel = [_HEADER]
    for no, kode, nama, metode, nominal, potongan, dibayar in _rows(draft):
        tabel.append(
            [
                str(no),
                kode,
                nama,
                metode,
                *(format_ribuan(Decimal(angka)) for angka in (nominal, potongan, dibayar)),
            ]
        )
    tabel.append(
        ["TOTAL", "", "", "", *(format_ribuan(Decimal(angka)) for angka in _totals(draft))]
    )
    gaya_tabel = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#006D44")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (4, 0), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F3F7F5")]),
        ]
    )
    isi.append(Table(tabel, repeatRows=1, style=gaya_tabel))
    stream = BytesIO()
    SimpleDocTemplate(stream, pagesize=A4, title=draft.nama).build(isi)
    return stream.getvalue()
