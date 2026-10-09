from decimal import Decimal
from io import BytesIO

from django.test import RequestFactory
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from api.models import BankSampah, DetailTransaksi, Nasabah, Transaksi, User
from apps.reporting.api import export_excel
from apps.waste_catalog.api import buat_jenis_sampah


class ExportEdgeCaseTests(APITestCase):
    def setUp(self) -> None:
        self.bank = BankSampah.objects.create(nama="Bank Laporan", alamat="Depok", kota="Depok")
        self.pengelola = User.objects.create_user(
            email="laporan@example.test",
            nama="Pengurus",
            bank_sampah=self.bank,
            is_profile_complete=True,
            is_primary_pengelola=True,
        )
        self.nasabah = Nasabah.objects.create(
            bank_sampah=self.bank, nomor="NAS-0001", nama="Budi", no_hp="+628111111111"
        )

    def test_exporting_nothing_yields_a_workbook_with_empty_riwayat(self) -> None:
        request = RequestFactory().get("/api/v1/transaksi/export")

        content, filename = export_excel(Transaksi.objects.none(), request)

        workbook = load_workbook(BytesIO(content))
        self.assertTrue(filename.endswith(".xlsx"))
        self.assertIn("Riwayat Transaksi", workbook.sheetnames)

    def test_unknown_period_is_labelled_as_all_periods(self) -> None:
        jenis = buat_jenis_sampah(
            bank_sampah=self.bank, nomor="PET", nama_sampah="Botol PET", harga_per_kg=Decimal(3000)
        )
        transaksi = Transaksi.objects.create(
            nasabah=self.nasabah, bank_sampah=self.bank, dicatat_oleh=self.pengelola
        )
        DetailTransaksi.objects.create(
            transaksi=transaksi,
            jenis_sampah=jenis,
            nama_sampah_snapshot="Botol PET",
            kategori_snapshot="plastik",
            berat=Decimal("2.000"),
            harga_snapshot=Decimal(3000),
            subtotal=Decimal(6000),
        )
        self.client.force_authenticate(self.pengelola)

        response = self.client.get("/api/v1/transaksi/export", {"periode": "semua"})

        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content))["Riwayat Transaksi"]
        cells = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value]
        self.assertTrue(any("Semua Periode" in value for value in cells), cells[:10])
