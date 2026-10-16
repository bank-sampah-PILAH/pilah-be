"""PIL-248 contract: the screen and XLSX share a filtered, bank-owned ledger."""

from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any, cast
from unittest.mock import patch
from uuid import uuid4

from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITransactionTestCase
from rest_framework_simplejwt.tokens import AccessToken

from api.models import BankSampah, DetailTransaksi, JenisSampah, Nasabah, Pencairan, Transaksi, User


class OperationalReportFixture(APITransactionTestCase):
    feed_url = "/api/v1/aktivitas"
    export_url = "/api/v1/transaksi/export"
    params = {"periode": "custom", "dari_tanggal": "2026-10-01", "sampai_tanggal": "2026-10-07"}

    def setUp(self) -> None:
        self.bank = BankSampah.objects.create(nama="Melati", no_hp_pic="081200")
        self.manager = User.objects.create_user(
            email="report@example.test",
            nama="Pengurus",
            role=User.Role.PENGELOLA,
            bank_sampah=self.bank,
            is_profile_complete=True,
        )
        self.member = Nasabah.objects.create(
            bank_sampah=self.bank,
            nomor="001",
            nama="Siti",
            no_hp="081201",
            alamat="Depok",
        )
        self.waste = JenisSampah.objects.create(
            bank_sampah=self.bank,
            nomor="PET",
            nama_sampah="Botol",
            harga_per_kg=8000,
        )
        self.first = self.deposit("2026-10-01T00:00:00+07:00", "12000", "1.500")
        self.second = self.deposit("2026-10-07T23:59:59+07:00", "8000", "2.000")
        self.payout = Pencairan.objects.create(
            bank_sampah=self.bank,
            nasabah=self.member,
            dicatat_oleh=self.manager,
            tanggal=datetime.fromisoformat("2026-10-04T12:00:00+07:00"),
            nominal=5000,
            metode="tunai",
            saldo_sebelum=12000,
            saldo_sesudah=7000,
        )
        self.deposit("2026-09-30T23:59:59+07:00", "6000", "1.000")
        self.client.force_authenticate(self.manager)

    def deposit(self, date: str, amount: str, weight: str) -> Transaksi:
        event = Transaksi.objects.create(
            bank_sampah=self.bank,
            nasabah=self.member,
            dicatat_oleh=self.manager,
            tanggal=datetime.fromisoformat(date),
            total_nilai=Decimal(amount),
        )
        DetailTransaksi.objects.create(
            transaksi=event,
            jenis_sampah=self.waste,
            nama_sampah_snapshot="Botol",
            kategori_snapshot="plastik",
            harga_snapshot=8000,
            berat=Decimal(weight),
            subtotal=Decimal(amount),
        )
        return event

    def summary(self, **filters: Any) -> dict[str, Any]:
        response = self.client.get(self.feed_url, {**self.params, **filters})
        self.assertEqual(response.status_code, 200, response.data)
        return cast(dict[str, Any], response.data["summary"])

    def workbook(self, **filters: Any) -> Any:
        response = self.client.get(self.export_url, {**self.params, **filters})
        self.assertEqual(response.status_code, 200, getattr(response, "data", None))
        return load_workbook(BytesIO(response.content), data_only=False)


class OperationalReportTests(OperationalReportFixture):
    def test_positive_summary_counts_events_not_detail_rows_or_page(self) -> None:
        DetailTransaksi.objects.create(
            transaksi=self.first,
            jenis_sampah=self.waste,
            nama_sampah_snapshot="Kardus",
            kategori_snapshot="kertas",
            harga_snapshot=8000,
            berat=Decimal("0.250"),
            subtotal=0,
        )
        expected = {
            "jumlah_transaksi": 3,
            "total_berat": "3.750",
            "total_setoran": 20000,
            "total_pencairan": 5000,
            "saldo": 15000,
        }
        self.assertEqual(self.summary(page_size=1), expected)
        self.assertEqual(self.summary(page_size=1, page=2), expected)

    def test_positive_export_has_operational_summary_and_both_types(self) -> None:
        wb = self.workbook(page=2, page_size=1)
        metrics = dict(wb["Ringkasan"].iter_rows(min_row=2, max_row=6, values_only=True))
        self.assertEqual(
            metrics,
            {
                "Jumlah transaksi": 3,
                "Total berat (kg)": 3.5,
                "Nilai setoran (Rp)": 20000,
                "Total pencairan (Rp)": 5000,
                "Saldo periode (Rp)": 15000,
            },
        )
        raw = list(wb["Riwayat Transaksi"].iter_rows(min_row=5, values_only=True))
        self.assertEqual(len(raw), 3)
        self.assertEqual([row[10] for row in raw], ["setoran", "pencairan", "setoran"])
        self.assertEqual(sum(row[12] or 0 for row in raw), 5000)

    def test_type_filters_match_feed_summary_and_export(self) -> None:
        for kind, count, value, payout, weight in (
            ("setoran", 2, 20000, 0, "3.500"),
            ("pencairan", 1, 0, 5000, "0.000"),
        ):
            with self.subTest(kind=kind):
                metrics = self.summary(tipe=kind)
                self.assertEqual(
                    metrics,
                    {
                        "jumlah_transaksi": count,
                        "total_berat": weight,
                        "total_setoran": value,
                        "total_pencairan": payout,
                        "saldo": value - payout,
                    },
                )
                rows = list(
                    self.workbook(tipe=kind)["Riwayat Transaksi"].iter_rows(
                        min_row=5, values_only=True
                    )
                )
                self.assertEqual(len(rows), count)
                self.assertTrue(all(row[10] == kind for row in rows))

    def test_search_matches_both_sources_and_no_match_has_zero_summary(self) -> None:
        self.assertEqual(self.summary(search="  siTI  ")["jumlah_transaksi"], 3)
        self.assertEqual(
            self.summary(search="absent"),
            {
                "jumlah_transaksi": 0,
                "total_berat": "0.000",
                "total_setoran": 0,
                "total_pencairan": 0,
                "saldo": 0,
            },
        )
        self.assertEqual(
            self.client.get(self.export_url, {**self.params, "search": "absent"}).status_code, 400
        )
        self.assertEqual(
            len(
                list(
                    self.workbook(search="S")["Riwayat Transaksi"].iter_rows(
                        min_row=5, values_only=True
                    )
                )
            ),
            3,
        )

    def test_custom_boundaries_and_single_day_are_inclusive(self) -> None:
        self.assertEqual(self.summary()["jumlah_transaksi"], 3)
        self.assertEqual(self.summary(dari_tanggal="2026-10-07")["total_setoran"], 8000)
        self.assertEqual(self.summary(sampai_tanggal="2026-10-01")["total_setoran"], 12000)

    def test_calendar_and_rolling_filters_match_export(self) -> None:
        fixed = datetime.fromisoformat("2026-10-07T12:00:00+07:00")
        with patch("django.utils.timezone.now", return_value=fixed):
            for period, count in (
                ("hari_ini", 1),
                ("bulan_ini", 3),
                ("bulan_lalu", 1),
                ("1_bulan", 4),
                ("semua", 4),
            ):
                with self.subTest(period=period):
                    self.assertEqual(self.summary(periode=period)["jumlah_transaksi"], count)
                    wb = self.workbook(periode=period)
                    self.assertEqual(wb["Ringkasan"]["B2"].value, count)

    def test_cross_bank_and_client_bank_override_do_not_leak(self) -> None:
        foreign_bank = BankSampah.objects.create(nama="Mawar", no_hp_pic="081202")
        foreign = Nasabah.objects.create(
            bank_sampah=foreign_bank, nomor="001", nama="Siti", no_hp="081203"
        )
        Transaksi.objects.create(
            bank_sampah=foreign_bank,
            nasabah=foreign,
            dicatat_oleh=self.manager,
            total_nilai=99999,
            tanggal=self.first.tanggal,
        )
        Pencairan.objects.create(
            bank_sampah=foreign_bank,
            nasabah=foreign,
            dicatat_oleh=self.manager,
            nominal=99999,
            metode="tunai",
            tanggal=self.first.tanggal,
            saldo_sebelum=99999,
            saldo_sesudah=0,
        )
        self.assertEqual(self.summary(bank_sampah_id=str(foreign_bank.id))["total_setoran"], 20000)
        self.assertEqual(
            self.workbook(bank_sampah_id=str(foreign_bank.id))["Ringkasan"]["B5"].value, 5000
        )

    def test_nasabah_superadmin_and_anonymous_are_denied(self) -> None:
        for role in (User.Role.NASABAH, User.Role.SUPERADMIN):
            self.manager.role = role
            self.manager.save(update_fields=["role"])
            for url in (self.feed_url, self.export_url):
                self.assertEqual(self.client.get(url, self.params).status_code, 403)
        self.client.force_authenticate(None)
        for url in (self.feed_url, self.export_url):
            self.assertEqual(self.client.get(url, self.params).status_code, 401)

    def test_inactive_bank_and_incomplete_manager_are_denied(self) -> None:
        self.bank.is_active = False
        self.bank.save(update_fields=["is_active"])
        for url in (self.feed_url, self.export_url):
            self.assertEqual(self.client.get(url, self.params).status_code, 403)
        self.bank.is_active = True
        self.bank.save(update_fields=["is_active"])
        self.manager.is_profile_complete = False
        self.manager.save(update_fields=["is_profile_complete"])
        for url in (self.feed_url, self.export_url):
            self.assertEqual(self.client.get(url, self.params).status_code, 403)

    def test_expired_and_tampered_tokens_are_denied(self) -> None:
        self.client.force_authenticate(None)
        token = AccessToken.for_user(self.manager)
        token.set_exp(from_time=timezone.now() - timedelta(days=2), lifetime=timedelta(seconds=1))
        for raw in (str(token), "invalid.jwt.signature"):
            self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {raw}")
            for url in (self.feed_url, self.export_url):
                self.assertEqual(self.client.get(url, self.params).status_code, 401)

    def test_invalid_filters_fail_closed_on_both_endpoints(self) -> None:
        for filters, status in (
            ({"periode": "unknown"}, 422),
            ({"tipe": "unknown"}, 422),
            ({"dari_tanggal": "bad"}, 422),
            ({"dari_tanggal": "2026-10-08"}, 400),
            ({"sampai_tanggal": ""}, 422),
        ):
            for url in (self.feed_url, self.export_url):
                with self.subTest(url=url, filters=filters):
                    self.assertEqual(
                        self.client.get(url, {**self.params, **filters}).status_code, status
                    )

    def test_money_is_rounded_per_event_without_float_or_join_inflation(self) -> None:
        self.first.total_nilai = Decimal("12000.99")
        self.first.save(update_fields=["total_nilai"])
        self.second.total_nilai = Decimal("8000.99")
        self.second.save(update_fields=["total_nilai"])
        self.payout.nominal = Decimal("5000.99")
        self.payout.save(update_fields=["nominal"])
        self.assertEqual(self.summary()["saldo"], 15000)
        self.assertEqual(self.workbook()["Ringkasan"]["B6"].value, 15000)

    def test_formula_like_user_text_is_not_an_executable_excel_formula(self) -> None:
        for payload in (
            '=HYPERLINK("https://example.invalid","x")',
            "+cmd",
            "-1+1",
            "@SUM(1)",
            "\t=1+1",
        ):
            with self.subTest(payload=payload):
                self.member.nama = payload
                self.member.save(update_fields=["nama"])
                cell = self.workbook()["Riwayat Transaksi"]["D5"]
                self.assertEqual(cell.data_type, "s")
                self.assertTrue(str(cell.value).startswith("'"))

    def test_sql_like_search_is_literal_and_not_a_filter_bypass(self) -> None:
        self.assertEqual(self.summary(search="' OR 1=1 --")["jumlah_transaksi"], 0)

    def test_financial_responses_are_private_and_not_cacheable(self) -> None:
        for url in (self.feed_url, self.export_url):
            response = self.client.get(url, self.params)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response["Cache-Control"], "private, no-store")
            self.assertEqual(response["X-Content-Type-Options"], "nosniff")

    def test_empty_bank_returns_zero_summary_and_export_error(self) -> None:
        Transaksi.objects.all().delete()
        Pencairan.objects.all().delete()
        self.assertEqual(self.summary()["jumlah_transaksi"], 0)
        response = self.client.get(self.export_url, self.params)
        self.assertEqual(response.status_code, 400)
        self.assertIn("Tidak ada data", response.data["error"])

    def test_payout_only_bank_can_export(self) -> None:
        Transaksi.objects.all().delete()
        self.assertEqual(self.summary()["saldo"], -5000)
        wb = self.workbook()
        self.assertEqual(wb["Ringkasan"]["B2"].value, 1)
        self.assertEqual(wb["Riwayat Transaksi"]["K5"].value, "pencairan")

    def test_payout_only_category_sheet_has_zero_not_a_circular_formula(self) -> None:
        wb = self.workbook(tipe="pencairan")
        self.assertEqual(wb["Laporan"]["E2"].value, 0)

    def test_export_metadata_names_the_actual_selected_period(self) -> None:
        for months in (1, 3, 6, 12):
            with self.subTest(months=months):
                label = self.workbook(periode=f"{months}_bulan")["Riwayat Transaksi"]["A2"].value
                self.assertIn(f"{months} Bulan Terakhir", label)
                self.assertNotIn("Semua Periode", label)
        with self.subTest(default="semua"):
            response = self.client.get(self.export_url)
            self.assertEqual(response.status_code, 200)
            label = load_workbook(BytesIO(response.content))["Riwayat Transaksi"]["A2"].value
            self.assertIn("Semua Periode", label)

    def test_member_filter_scopes_both_ledgers_summary_and_category_sheet(self) -> None:
        other = Nasabah.objects.create(
            bank_sampah=self.bank, nomor="002", nama="Ani", no_hp="081299", email="ani@example.test"
        )
        deposit = self.deposit("2026-10-03T12:00:00+07:00", "9000", "1.125")
        deposit.nasabah = other
        deposit.save(update_fields=["nasabah"])
        payout = Pencairan.objects.create(
            bank_sampah=self.bank,
            nasabah=other,
            dicatat_oleh=self.manager,
            tanggal=self.payout.tanggal,
            nominal=4000,
            metode="tunai",
            saldo_sebelum=9000,
            saldo_sesudah=5000,
        )
        for member, ids, value, paid, weight in [
            (
                self.member,
                {str(self.first.id), str(self.second.id), str(self.payout.id)},
                20000,
                5000,
                3.5,
            ),
            (other, {str(deposit.id), str(payout.id)}, 9000, 4000, 1.125),
        ]:
            with self.subTest(member=member.nomor):
                wb = self.workbook(nasabah_id=str(member.id))
                rows = list(wb["Riwayat Transaksi"].iter_rows(min_row=5, values_only=True))
                self.assertEqual({row[13] for row in rows}, ids)
                self.assertTrue(all(row[4] == member.nomor for row in rows))
                self.assertEqual(wb["Ringkasan"]["B2"].value, len(ids))
                self.assertEqual(wb["Ringkasan"]["B3"].value, weight)
                self.assertEqual(wb["Ringkasan"]["B4"].value, value)
                self.assertEqual(wb["Ringkasan"]["B5"].value, paid)
                self.assertEqual(wb["Ringkasan"]["B6"].value, value - paid)
                self.assertEqual(
                    sum(
                        row[3] or 0
                        for row in wb["Laporan"].iter_rows(
                            min_row=2, max_row=wb["Laporan"].max_row - 1, values_only=True
                        )
                    ),
                    weight,
                )
                self.assertEqual(
                    self.summary(nasabah_id=str(member.id))["jumlah_transaksi"], len(ids)
                )
                self.assertEqual(
                    self.workbook(nasabah_id=str(member.id), tipe="pencairan")["Ringkasan"][
                        "B4"
                    ].value,
                    0,
                )

    def test_member_filter_rejects_invalid_ids_and_does_not_expand_bank_scope(self) -> None:
        foreign_bank = BankSampah.objects.create(nama="Other", no_hp_pic="081244")
        foreign = Nasabah.objects.create(
            bank_sampah=foreign_bank, nomor="X", nama="Other", no_hp="081255"
        )
        for member_id in [str(foreign.id), str(uuid4())]:
            with self.subTest(member_id=member_id):
                response = self.client.get(
                    self.export_url, {**self.params, "nasabah_id": member_id}
                )
                self.assertEqual(response.status_code, 400)
        response = self.client.get(self.export_url, {**self.params, "nasabah_id": "invalid"})
        self.assertEqual(response.status_code, 422)
