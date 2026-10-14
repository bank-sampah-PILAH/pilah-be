from calendar import monthrange
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any, cast
from uuid import UUID

from django.db.models import QuerySet, Sum
from django.http import HttpRequest
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from api.models import DetailTransaksi, Pencairan, Transaksi
from apps.reporting.operational import OperationalSummary, operational_summary
from shared_kernel.kalkulasi import bulatkan_rupiah


def export_excel(
    queryset: QuerySet[Transaksi],
    request: HttpRequest | None = None,
    *,
    payouts: QuerySet[Pencairan] | None = None,
    summary: OperationalSummary | None = None,
) -> tuple[bytes, str]:
    if payouts is None:
        payouts = Pencairan.objects.none()
    if summary is None:
        summary = operational_summary(queryset, payouts)
    wb = Workbook()
    _fill_summary_export_sheet(cast(Worksheet, wb.active), queryset)
    _fill_raw_export_sheet(wb.create_sheet("Riwayat Transaksi"), queryset, request, payouts)
    ws = wb.create_sheet("Ringkasan")
    ws.append(["Metrik", "Nilai"])
    ws.append(["Jumlah transaksi", summary["jumlah_transaksi"]])
    ws.append(["Total berat (kg)", float(summary["total_berat"])])
    ws.append(["Nilai setoran (Rp)", summary["total_setoran"]])
    ws.append(["Total pencairan (Rp)", summary["total_pencairan"]])
    ws.append(["Saldo periode (Rp)", summary["saldo"]])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 25
    ws.freeze_panes = "A2"
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue(), f"PILAH_Laporan_{_month_label()}.xlsx"


def _month_label() -> str:
    return timezone.localdate().strftime("%B_%Y")


def _excel_text(value: str) -> str:
    """User-controlled text must never become a spreadsheet expression."""
    if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
        return "'" + value
    return value


def _fill_summary_export_sheet(ws: Worksheet, queryset: QuerySet[Transaksi]) -> None:
    ws.title = "Laporan"
    ws.append(["Jenis", "Sampah", "Harga per kg", "Jumlah kg", "Total"])
    rows = (
        DetailTransaksi.objects.filter(transaksi__in=queryset)
        .values("kategori_snapshot", "nama_sampah_snapshot", "harga_snapshot")
        .annotate(total_kg=Sum("berat"))
        .order_by("kategori_snapshot", "nama_sampah_snapshot")
    )
    current_row = 2
    for row in rows:
        ws.append(
            [
                _excel_text(row["kategori_snapshot"]),
                _excel_text(row["nama_sampah_snapshot"]),
                float(row["harga_snapshot"]),
                float(row["total_kg"]),
                f"=C{current_row}*D{current_row}",
            ]
        )
        current_row += 1
    ws.append(["", "", "", "TOTAL", f"=SUM(E2:E{current_row - 1})" if current_row > 2 else 0])
    for column in ["A", "B", "C", "D", "E"]:
        ws.column_dimensions[column].width = 22


def _fill_raw_export_sheet(
    ws: Worksheet,
    queryset: QuerySet[Transaksi],
    request: HttpRequest | None = None,
    payouts: QuerySet[Pencairan] | None = None,
) -> None:
    ws.title = "Riwayat Transaksi"
    ws.merge_cells("A1:N1")
    ws.merge_cells("A2:N2")
    ws["A1"] = "PILAH - Riwayat Transaksi"
    ws["A2"] = _export_filter_label(queryset, request)
    ws.append([])
    ws.append(
        [
            "No",
            "Tanggal",
            "Waktu",
            "Nama Nasabah",
            "ID Nasabah",
            "Jenis Sampah",
            "Berat (kg)",
            "Harga/kg (Rp)",
            "Subtotal (Rp)",
            "Saldo Setelah Transaksi (Rp)",
            "Tipe",
            "Nilai Setoran (Rp)",
            "Pencairan (Rp)",
            "ID Aktivitas",
        ]
    )

    saldo_after_by_transaction = _saldo_after_by_transaction(queryset)
    events: list[Transaksi | Pencairan] = list(
        queryset.select_related("nasabah").prefetch_related("items")
    )
    if payouts is not None:
        events.extend(payouts.select_related("nasabah"))
    events.sort(
        key=lambda event: (
            event.tanggal,
            event.created_at,
            "setoran" if isinstance(event, Transaksi) else "pencairan",
            str(event.id),
        ),
        reverse=True,
    )
    number = 0
    for event in events:
        local_datetime = timezone.localtime(event.tanggal)
        common: list[Any] = [
            local_datetime.strftime("%d/%m/%Y"),
            local_datetime.strftime("%H:%M"),
            _excel_text(event.nasabah.nama),
            _excel_text(event.nasabah.nomor),
        ]
        if isinstance(event, Pencairan):
            number += 1
            ws.append(
                [
                    number,
                    *common,
                    "Pencairan",
                    0,
                    0,
                    0,
                    int(event.saldo_sesudah),
                    "pencairan",
                    0,
                    int(event.nominal),
                    str(event.id),
                ]
            )
            continue
        items = sorted(event.items.all(), key=lambda item: str(item.id))
        if not items:
            number += 1
            ws.append(
                [
                    number,
                    *common,
                    "Setoran",
                    0,
                    0,
                    int(event.total_nilai),
                    int(saldo_after_by_transaction[event.id]),
                    "setoran",
                    int(event.total_nilai),
                    0,
                    str(event.id),
                ]
            )
        for index, item in enumerate(items):
            number += 1
            ws.append(
                [
                    number,
                    *common,
                    _excel_text(item.nama_sampah_snapshot),
                    float(item.berat),
                    float(item.harga_snapshot),
                    int(item.subtotal),
                    int(saldo_after_by_transaction[event.id]),
                    "setoran",
                    int(event.total_nilai) if index == 0 else 0,
                    0,
                    str(event.id),
                ]
            )

    _style_raw_export_sheet(ws)


def _saldo_after_by_transaction(queryset: QuerySet[Transaksi]) -> dict[UUID, Decimal]:
    target_transactions = list(queryset.select_related("bank_sampah").order_by("tanggal", "id"))
    if not target_transactions:
        return {}

    target_ids = {trans.id for trans in target_transactions}
    nasabah_ids = {trans.nasabah_id for trans in target_transactions}
    latest_transaction = target_transactions[-1]
    bank_sampah = latest_transaction.bank_sampah
    running_balances = {nasabah_id: Decimal("0.00") for nasabah_id in nasabah_ids}
    saldo_after = {}

    transactions = (
        Transaksi.objects.filter(
            bank_sampah=bank_sampah,
            nasabah_id__in=nasabah_ids,
            tanggal__lte=latest_transaction.tanggal,
        )
        .only("id", "nasabah_id", "total_nilai", "tanggal")
        .order_by("nasabah_id", "tanggal", "id")
    )
    pencairan = (
        Pencairan.objects.filter(
            bank_sampah=bank_sampah,
            nasabah_id__in=nasabah_ids,
            tanggal__lte=latest_transaction.tanggal,
        )
        .only("id", "nasabah_id", "saldo_sebelum", "saldo_sesudah", "tanggal")
        .order_by("nasabah_id", "tanggal", "id")
    )
    # Merge both ledgers so a setoran recorded after a pencairan reports the
    # balance that actually remains.
    events: list[tuple[UUID, datetime, UUID, Decimal, bool]] = [
        (trans.nasabah_id, trans.tanggal, trans.id, trans.total_nilai, True)
        for trans in transactions
    ]
    events.extend(
        (cair.nasabah_id, cair.tanggal, cair.id, cair.saldo_sesudah - cair.saldo_sebelum, False)
        for cair in pencairan
    )
    events.sort(key=lambda event: (str(event[0]), event[1], not event[4], str(event[2])))
    for nasabah_id, _tanggal, event_id, delta, is_transaksi in events:
        running_balances[nasabah_id] += delta
        if is_transaksi and event_id in target_ids:
            # History matches the stored saldo, which is rounded down to full Rupiah.
            saldo_after[event_id] = bulatkan_rupiah(running_balances[nasabah_id])

    return saldo_after


def _export_filter_label(queryset: QuerySet[Transaksi], request: HttpRequest | None) -> str:
    first_transaction = queryset.first()
    bank_name = first_transaction.bank_sampah.nama if first_transaction else "-"
    downloaded_at = timezone.localtime().strftime("%d/%m/%Y %H:%M")
    periode_text = "Semua Periode"
    if request:
        periode_text = period_label(request)
    kind = request.GET.get("tipe", "semua") if request else "semua"
    search = request.GET.get("search", "").strip() if request else ""
    user_bank = getattr(getattr(request, "user", None), "bank_sampah", None)
    if not first_transaction and user_bank is not None:
        bank_name = user_bank.nama
    return (
        f"Filter periode: {periode_text}   |   Bank Sampah: {bank_name}   |   Diunduh: {downloaded_at}"
        f"   |   Tipe: {kind}   |   Pencarian: {search or '-'}"
    )


def period_label(request: HttpRequest) -> str:
    """Human label for the ``periode`` param, shared by all report exports."""
    periode = request.GET.get("periode", "hari_ini")
    today = timezone.localdate()
    if periode == "hari_ini":
        return f"Hari Ini ({today.strftime('%d/%m/%Y')})"
    if periode == "minggu_ini":
        start = today - timedelta(days=today.weekday())
        return f"Minggu Ini ({start.strftime('%d/%m/%Y')} - {today.strftime('%d/%m/%Y')})"
    if periode == "bulan_ini":
        end_day = monthrange(today.year, today.month)[1]
        return f"Bulan Ini (1 - {end_day} {_month_name(today.month)} {today.year})"
    if periode == "bulan_lalu":
        first_this_month = today.replace(day=1)
        previous_month = first_this_month - timedelta(days=1)
        end_day = monthrange(previous_month.year, previous_month.month)[1]
        return (
            f"Bulan Lalu (1 - {end_day} {_month_name(previous_month.month)} {previous_month.year})"
        )
    if periode == "custom":
        start_date = request.GET.get("dari_tanggal", "")
        end_date = request.GET.get("sampai_tanggal", "")
        return f"Custom ({start_date} - {end_date})"
    return "Semua Periode"


def _month_name(month: int) -> str:
    return [
        "Januari",
        "Februari",
        "Maret",
        "April",
        "Mei",
        "Juni",
        "Juli",
        "Agustus",
        "September",
        "Oktober",
        "November",
        "Desember",
    ][month - 1]


def _style_raw_export_sheet(ws: Worksheet) -> None:
    ws.freeze_panes = "A5"
    widths = {
        "A": 5,
        "B": 13,
        "C": 9,
        "D": 18,
        "E": 13,
        "F": 16,
        "G": 11,
        "H": 13,
        "I": 13,
        "J": 28,
        "K": 16,
        "L": 24,
        "M": 24,
        "N": 38,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    title_fill = PatternFill("solid", fgColor="166534")
    subtitle_fill = PatternFill("solid", fgColor="DCFCE7")
    header_fill = PatternFill("solid", fgColor="16A34A")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A2"].font = Font(size=10, color="374151")
    ws["A2"].fill = subtitle_fill
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 22

    for row in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=10):
        for cell in row:
            cell.fill = title_fill if cell.row == 1 else subtitle_fill

    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=10):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
        row[0].alignment = Alignment(horizontal="center", vertical="center")
        row[6].number_format = "0.##"
        for cell in row[7:10]:
            cell.number_format = "#,##0"
