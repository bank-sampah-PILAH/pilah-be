"""PIL-428: a voided setoran counts toward no laporan, export or dashboard total."""

from decimal import Decimal
from io import BytesIO

from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.request import Request
from rest_framework.test import APIRequestFactory, APITestCase, APITransactionTestCase

from api.models import BankSampah, Nasabah, Transaksi, User
from apps.reporting.statement import build_statement
from apps.waste_catalog.api import buat_jenis_sampah


class _VoidedSetoranFixture(APITransactionTestCase):
    """One tercatat setoran of 3.000 and one voided of 6.000, for one nasabah."""

    def setUp(self) -> None:
        self.bank = BankSampah.objects.create(
            nama="Bank Laporan", alamat="Depok", kota="Depok", no_hp_pic="+628123456789"
        )
        self.pengelola = User.objects.create_user(
            email="laporan@example.test",
            nama="Pengelola",
            bank_sampah=self.bank,
            is_profile_complete=True,
            is_primary_pengelola=True,
        )
        self.nasabah = Nasabah.objects.create(
            bank_sampah=self.bank,
            nomor="NAS-0001",
            nama="Budi",
            no_hp="+628111111111",
            alamat="Jl. A",
        )
        self.pet = buat_jenis_sampah(
            bank_sampah=self.bank, nomor="PET", nama_sampah="Botol PET", harga_per_kg=Decimal(3000)
        )
        self.client.force_authenticate(self.pengelola)
        self.tetap = self._setor("1.000")
        self.batal = self._setor("2.000")
        response = self.client.post(
            f"/api/v1/transaksi/{self.batal.id}/batalkan", {"alasan": "Dobel"}, format="json"
        )
        self.assertEqual(response.status_code, 200)

    def _setor(self, berat: str) -> Transaksi:
        response = self.client.post(
            "/api/v1/transaksi",
            {
                "nasabah_id": str(self.nasabah.id),
                "items": [{"jenis_sampah_id": str(self.pet.id), "berat": berat}],
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201)
        return Transaksi.objects.get(pk=response.data["id"])


class VoidedSetoranReportingTests(_VoidedSetoranFixture, APITestCase):
    def test_dashboard_stats_leave_the_voided_setoran_out(self) -> None:
        response = self.client.get("/api/v1/dashboard/stats")

        self.assertEqual(response.data["transaksi_bulan_ini"], 1)
        self.assertEqual(Decimal(response.data["total_nilai_bulan_ini"]), Decimal(3000))
        self.assertEqual(Decimal(response.data["total_sampah_kg_bulan_ini"]), Decimal(1))

    def test_dashboard_trends_and_categories_leave_the_voided_setoran_out(self) -> None:
        hari_ini = timezone.localdate().isoformat()
        periode = {"start_date": hari_ini, "end_date": hari_ini}

        trends = self.client.get(
            "/api/v1/dashboard/operational-trends", {**periode, "granularity": "day"}
        )
        categories = self.client.get("/api/v1/dashboard/waste-categories", periode)

        bucket = trends.data["buckets"][0]
        self.assertEqual((bucket["setoran_total_nilai"], bucket["setoran_count"]), ("3000.00", 1))
        category = categories.data["categories"][0]
        self.assertEqual(category["deposit_count"], 1)
        self.assertEqual(Decimal(category["total_weight_kg"]), Decimal(1))

    def test_recent_transactions_show_the_void(self) -> None:
        response = self.client.get("/api/v1/dashboard/recent-transactions")

        statuses = {row["id"]: row["status"] for row in response.data["transactions"]}
        self.assertEqual(statuses[str(self.batal.id)], "dibatalkan")

    def test_nasabah_statement_leaves_the_voided_setoran_out(self) -> None:
        request = APIRequestFactory().get("/api/v1/nasabah/me/riwayat/export-pdf")

        data = build_statement(self.nasabah, Request(request))

        self.assertEqual([row.id for row in data.rows], [self.tetap.id])
        self.assertEqual(data.rows[-1].saldo, Decimal(3000))


class VoidedSetoranExportTests(_VoidedSetoranFixture):
    # The export opens its own outermost REPEATABLE READ transaction.

    def test_excel_export_leaves_the_voided_setoran_out(self) -> None:
        response = self.client.get("/api/v1/transaksi/export", {"periode": "semua"})

        wb = load_workbook(BytesIO(response.content))
        metrics = dict(wb["Ringkasan"].iter_rows(min_row=2, max_row=6, values_only=True))
        self.assertEqual(metrics["Nilai setoran (Rp)"], 3000)
        self.assertEqual(metrics["Jumlah transaksi"], 1)
        raw = list(wb["Riwayat Transaksi"].iter_rows(min_row=5, values_only=True))
        self.assertNotIn(str(self.batal.id), {cell for row in raw for cell in row})
        laporan = list(wb["Laporan"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(laporan[-1][-1], 3000)

    def test_export_of_only_voided_setoran_has_no_data(self) -> None:
        Transaksi.objects.filter(pk=self.tetap.pk).update(
            status=Transaksi.Status.DIBATALKAN,
            alasan_batal="x",
            dibatalkan_oleh=self.pengelola,
            dibatalkan_pada=timezone.now(),
        )

        response = self.client.get("/api/v1/transaksi/export", {"periode": "semua"})

        self.assertEqual(response.status_code, 400)
